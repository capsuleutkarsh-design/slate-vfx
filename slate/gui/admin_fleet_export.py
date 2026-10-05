"""
Fleet report export helpers for Admin Panel.
"""

from datetime import datetime

from slate.core.domain.table_export import neutralise


def export_fleet_xlsx(output_path, records, summary, skipped, columns=None, header_for=None):
    """
    Create a colour-coded, bordered Excel workbook for the fleet report.

    columns / header_for are the report's own column order and human headers
    (admin_fleet_report_service), so the workbook and the CSV read the same.
    """
    from openpyxl import Workbook
    from openpyxl.styles import (
        PatternFill, Font, Alignment, Border, Side
    )
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    # Palette
    col_header_bg = "1E2A3A"
    col_header_fg = "FFFFFF"
    col_online_bg = "D4EDDA"
    col_online_fg = "155724"
    col_idle_bg = "FFF3CD"
    col_idle_fg = "856404"
    col_offline_bg = "F8D7DA"
    col_offline_fg = "721C24"
    col_unknown_bg = "E2E3E5"
    col_unknown_fg = "383D41"
    col_crit_bg = "FF0000"
    col_warn_bg = "FFA500"
    col_ok_bg = "28A745"
    col_summary_bg = "2C3E50"
    col_tile_online = "27AE60"
    col_tile_idle = "F39C12"
    col_tile_off = "E74C3C"

    thin = Side(style="thin", color="CCCCCC")
    Side(style="medium", color="999999")
    thick = Side(style="thick", color=col_header_bg)
    border_thin = Border(left=thin, right=thin, top=thin, bottom=thin)

    def make_fill(hex_color):
        return PatternFill("solid", fgColor=hex_color)

    def make_font(hex_color, bold=False, size=10):
        return Font(color=hex_color, bold=bold, size=size, name="Segoe UI")

    # Sheet 1: Summary
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_sum.sheet_view.showGridLines = False

    ws_sum.merge_cells("A1:F1")
    title_cell = ws_sum["A1"]
    title_cell.value = f"Slate Fleet Report - {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    title_cell.font = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")
    title_cell.fill = make_fill(col_summary_bg)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[1].height = 36

    # The tiles add up to the total: Unknown (no last report time) used to
    # be counted in the total and in no tile.
    not_responding = summary.get("not_responding", summary.get("idle", 0))
    tiles = [
        ("Machines", len(records), "3498DB"),
        ("Online", summary["online"], col_tile_online),
        ("Not responding", not_responding, col_tile_idle),
        ("Offline", summary["offline"], col_tile_off),
        ("Unknown", summary.get("unknown", 0), "95A5A6"),
        ("Unreadable files", skipped, "7F8C8D"),
    ]
    headers_row, values_row = 3, 4
    for col_idx, (label, val, color) in enumerate(tiles, start=1):
        h_cell = ws_sum.cell(row=headers_row, column=col_idx, value=label)
        h_cell.fill = make_fill(color)
        h_cell.font = make_font("FFFFFF", bold=True, size=9)
        h_cell.alignment = Alignment(horizontal="center", vertical="center")
        h_cell.border = border_thin
        ws_sum.row_dimensions[headers_row].height = 20

        v_cell = ws_sum.cell(row=values_row, column=col_idx, value=val)
        v_cell.fill = make_fill("F8F9FA")
        v_cell.font = Font(name="Segoe UI", size=22, bold=True, color=color)
        v_cell.alignment = Alignment(horizontal="center", vertical="center")
        v_cell.border = border_thin
        ws_sum.row_dimensions[values_row].height = 44
        ws_sum.column_dimensions[get_column_letter(col_idx)].width = 22

    # Sheet 2: Fleet data
    ws = wb.create_sheet("Fleet Data")
    ws.sheet_view.showGridLines = False

    if not records:
        ws["A1"] = "No data found."
        wb.save(output_path)
        return

    all_keys = list(columns) if columns else list(dict.fromkeys(k for r in records for k in r))
    if header_for is None:
        def header_for(key):
            return key.replace("_", " ").capitalize()

    header_fill = make_fill(col_header_bg)
    header_font = make_font(col_header_fg, bold=True, size=10)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_border = Border(left=thin, right=thin, top=thick, bottom=thick)

    for col_idx, key in enumerate(all_keys, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header_for(key))
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = header_border
    ws.row_dimensions[1].height = 30

    status_styles = {
        "Online": (col_online_bg, col_online_fg),
        "Not responding": (col_idle_bg, col_idle_fg),
        "Offline": (col_offline_bg, col_offline_fg),
        "Unknown": (col_unknown_bg, col_unknown_fg),
    }
    alert_fills = {
        "CRITICAL": make_fill(col_crit_bg),
        "WARNING": make_fill(col_warn_bg),
        "OK": make_fill(col_ok_bg),
    }
    alert_fonts = {
        "CRITICAL": make_font("FFFFFF", bold=True),
        "WARNING": make_font("000000", bold=True),
        "OK": make_font("FFFFFF", bold=True),
    }

    # No divider rows between the groups: they sat inside the filtered range
    # and got sorted in with the data. The records arrive grouped by status
    # and every row is coloured by it, which is grouping enough.
    row_idx = 2
    for record in records:
        status = record.get("status", "Unknown")

        row_fill = make_fill(status_styles.get(status, (col_unknown_bg, col_unknown_fg))[0])
        row_font = make_font(status_styles.get(status, (col_unknown_bg, col_unknown_fg))[1])

        for col_idx, key in enumerate(all_keys, start=1):
            val = record.get(key, "")
            # A PC report is written by the PC: its text must not become a formula (SYS2-004).
            cell = ws.cell(row=row_idx, column=col_idx, value=neutralise(val))
            cell.fill = row_fill
            cell.font = row_font
            cell.border = border_thin
            cell.alignment = Alignment(vertical="center")

            if key == "status":
                cell.font = make_font(status_styles.get(status, (col_unknown_bg, col_unknown_fg))[1], bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")

            if key.endswith("_usage_pct") and isinstance(val, (int, float)):
                cell.number_format = '0.0"%"'

            if key.endswith("_alert") and val in alert_fills:
                cell.fill = alert_fills[val]
                cell.font = alert_fonts[val]
                cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[row_idx].height = 18
        row_idx += 1

    for col_idx, key in enumerate(all_keys, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = max(
            len(str(header_for(key))),
            max((len(str(r.get(key, ""))) for r in records), default=0),
        )
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 40)

    ws.freeze_panes = "B2"
    # The whole table, not the header row only: Excel builds the filter
    # lists from this range.
    ws.auto_filter.ref = ws.dimensions

    wb.save(output_path)
