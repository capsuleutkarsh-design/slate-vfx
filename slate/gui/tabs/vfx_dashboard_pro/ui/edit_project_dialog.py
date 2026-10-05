"""
Edit a project: its name, root folder and Excel backup.

Labels used to collide ("Excel Rows:" over "Header Row:"), paths were cut off
from the left with no way to read them, and the advanced fields showed while
their box was unticked. The form is one column of labelled rows now, paths
show their start with the whole path as the tooltip, and the sheet layout
rows only appear when asked for. OK saves to the database (ProjectManager) -
never to the install folder.
"""

import os

from openpyxl import load_workbook
from PySide6.QtWidgets import (
    QCheckBox, QComboBox, QDialog, QDialogButtonBox, QDoubleSpinBox, QFileDialog, QHBoxLayout,
    QLabel, QLineEdit, QSpinBox, QVBoxLayout, QWidget,
)

from slate.core.infra.gate import Gate
from slate.gui.core.controls import form_layout, make_button, style_button


class _PathField(QWidget):
    """A path box with Browse; the whole path is the tooltip."""

    def __init__(self, value, browse, parent=None):
        super().__init__(parent)
        row = QHBoxLayout(self)
        row.setContentsMargins(0, 0, 0, 0)
        self.edit = QLineEdit(value or "")
        self.edit.textChanged.connect(lambda text: self.edit.setToolTip(text))
        self.edit.setToolTip(value or "")
        self.edit.setCursorPosition(0)
        row.addWidget(self.edit, 1)
        row.addWidget(make_button("Browse…", "secondary", on_click=browse))

    def text(self):
        return self.edit.text()

    def setText(self, value):
        self.edit.setText(value)
        self.edit.setCursorPosition(0)


class EditProjectDialog(QDialog):
    def __init__(self, project_config, parent=None):
        super().__init__(parent)
        self.setWindowTitle(f"Edit project - {project_config.code}")
        self.setMinimumWidth(600)
        self.project_config = project_config

        self.main_layout = QVBoxLayout(self)
        self.form = form_layout()

        self.code_input = QLineEdit(project_config.code)
        self.code_input.setReadOnly(True)
        self.code_input.setToolTip("A project's code cannot be changed: shots and history are filed under it.")
        self.code_input.setStyleSheet(f"color: {Gate.TEXT_DIM};")
        self.form.addRow("Code", self.code_input)

        self.name_input = QLineEdit(project_config.name)
        self.form.addRow("Name", self.name_input)

        self.folder_field = _PathField(project_config.folder_base, self.browse_folder)
        self.folder_input = self.folder_field.edit
        self.form.addRow("Project root", self.folder_field)

        self.excel_field = _PathField(project_config.excel_path, self.browse_excel)
        self.excel_input = self.excel_field.edit
        self.form.addRow("Excel backup", self.excel_field)

        self.sheet_combo = QComboBox()
        self.sheet_combo.setEditable(True)
        self.sheet_combo.setPlaceholderText("Sheet name")
        if project_config.sheet_name:
            self.sheet_combo.addItem(project_config.sheet_name)
            self.sheet_combo.setCurrentText(project_config.sheet_name)
        sheet_row = QWidget()
        sheet_layout = QHBoxLayout(sheet_row)
        sheet_layout.setContentsMargins(0, 0, 0, 0)
        sheet_layout.addWidget(self.sheet_combo, 1)
        self.sheet_btn = make_button("Read sheets", "secondary",
                                     tooltip="List the sheets in the Excel file", on_click=self.analyze_excel)
        sheet_layout.addWidget(self.sheet_btn)
        self.form.addRow("Sheet", sheet_row)

        # The rate image sequences play at; a movie keeps its own.
        self.fps_spin = QDoubleSpinBox()
        self.fps_spin.setRange(1.0, 120.0)
        self.fps_spin.setDecimals(3)
        self.fps_spin.setValue(float(getattr(project_config, "fps", 0) or 24.0))
        self.fps_spin.setToolTip("Used for image sequences in the lineup, EDLs, proxies and the "
                                 "player. Movies keep their own frame rate.")
        self.form.addRow("Frame rate (sequences)", self.fps_spin)

        self.adv_toggle = QCheckBox("Change the sheet layout (header and first data row)")
        self.adv_group = self.adv_toggle
        self.form.addRow("", self.adv_toggle)
        self.header_row_spin = QSpinBox()
        self.header_row_spin.setRange(1, 100)
        self.header_row_spin.setValue(project_config.header_row)
        self.form.addRow("Header row", self.header_row_spin)
        self.data_row_spin = QSpinBox()
        self.data_row_spin.setRange(1, 100)
        self.data_row_spin.setValue(project_config.data_start_row)
        self.form.addRow("First data row", self.data_row_spin)
        self.adv_toggle.toggled.connect(self._show_advanced)
        self._show_advanced(False)

        self.main_layout.addLayout(self.form)
        self.message = QLabel("")
        self.message.setWordWrap(True)
        self.message.setStyleSheet(f"color: {Gate.TEXT_2};")
        self.main_layout.addWidget(self.message)

        self.buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        ok = self.buttons.button(QDialogButtonBox.StandardButton.Ok)
        ok.setText("Save")
        style_button(ok, "primary")
        style_button(self.buttons.button(QDialogButtonBox.StandardButton.Cancel), "secondary")
        self.buttons.accepted.connect(self.validate_and_accept)
        self.buttons.rejected.connect(self.reject)
        self.main_layout.addWidget(self.buttons)

    def _show_advanced(self, shown: bool):
        for widget in (self.header_row_spin, self.data_row_spin):
            widget.setVisible(bool(shown))
            label = self.form.labelForField(widget)
            if label is not None:
                label.setVisible(bool(shown))

    def browse_excel(self):
        # openpyxl reads .xlsx/.xlsm only; an old .xls must be saved as .xlsx first.
        path, _ = QFileDialog.getOpenFileName(self, "Excel backup file", self.excel_input.text(),
                                              "Excel workbooks (*.xlsx *.xlsm)")
        if path:
            self.excel_field.setText(path)
            self.analyze_excel()

    def analyze_excel(self):
        path = self.excel_input.text().strip()
        if not path or not os.path.exists(path):
            self.message.setText("Pick an Excel file that exists first.")
            return
        try:
            wb = load_workbook(path, read_only=True, keep_vba=False)
            sheets = wb.sheetnames
            wb.close()
        except Exception as e:
            self.message.setText(f"The Excel file could not be read: {e}")
            return
        current = self.sheet_combo.currentText()
        self.sheet_combo.clear()
        self.sheet_combo.addItems(sheets)
        if current in sheets:
            self.sheet_combo.setCurrentText(current)
        else:
            best = next((sh for sh in sheets if "MASTER" in sh.upper()), None) or \
                next((sh for sh in sheets if "TRACKER" in sh.upper()), None)
            if best:
                self.sheet_combo.setCurrentText(best)
        self.message.setText(f"{len(sheets)} sheet{'s' if len(sheets) != 1 else ''} in the file.")

    def browse_folder(self):
        path = QFileDialog.getExistingDirectory(self, "Project root folder", self.folder_input.text())
        if path:
            self.folder_field.setText(path)

    def _refuse(self, text, widget):
        self.message.setText(text)
        self.message.setStyleSheet(f"color: {Gate.BAD};")
        widget.setFocus()

    def validate_and_accept(self):
        if not self.name_input.text().strip():
            self._refuse("The project needs a name.", self.name_input)
            return
        if self.data_row_spin.value() <= self.header_row_spin.value():
            # Data starting at or above the headings would be written over them.
            self.adv_toggle.setChecked(True)
            self._refuse("The first data row must be below the header row.", self.data_row_spin)
            return
        if self.excel_input.text().strip() and not self.sheet_combo.currentText().strip():
            self._refuse("Name the sheet the backup is written to.", self.sheet_combo)
            return
        self.accept()

    def get_data(self):
        return {
            "code": self.project_config.code,
            "name": self.name_input.text().strip(),
            "excel_path": self.excel_input.text().strip(),
            "folder_base": self.folder_input.text().strip(),
            "sheet_name": self.sheet_combo.currentText().strip(),
            "header_row": self.header_row_spin.value(),
            "data_start_row": self.data_row_spin.value(),
            "fps": self.fps_spin.value(),
        }
