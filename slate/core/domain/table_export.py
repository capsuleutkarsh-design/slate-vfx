"""
Tables out to a file - CSV or XLSX - safely.

No IT table could be exported at all, and the exports other screens had were
one-offs; one of them wrote cell text straight into CSV, so a name typed as
=HYPERLINK("http://...") or +cmd|' /C calc'!A0 became a live formula the moment
HR opened the file in Excel. This is the one way to do it:

    from slate.core.domain.table_export import export_rows
    export_rows("licences.csv", ["Software", "Seats"], [["Nuke", 12], ...])

    # from a QTableWidget / QTableView, exactly what is on screen:
    headers, rows = rows_from_qtable(table)

What it guarantees:

  * Formula injection is neutralised. In CSV, any text cell starting with
    = + - @ (or a tab / carriage return that Excel strips before looking) is
    written with a leading apostrophe, which Excel shows as text. In XLSX the
    cell is stored as text instead (put / append_row), with no apostrophe to
    show. Real numbers stay numbers - "-5" typed as text is quoted, -5 as a
    number is not.
  * CSV is UTF-8 with a byte-order mark, so Excel opens Hindi names, ₹ and
    emoji correctly instead of as mojibake.
  * XLSX (when openpyxl is available) keeps numbers and dates as real
    numbers and dates.
  * Only visible columns and rows are exported, in the order on screen.
"""

from __future__ import annotations

import csv
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Iterable, List, Sequence, Tuple

_DANGEROUS_START = ("=", "+", "-", "@", "\t", "\r", "\n", "＝", "＋", "－", "＠")
_SAFE_FILENAME = re.compile(r"[^A-Za-z0-9._ -]+")


def neutralise(value):
    """
    The value, made safe to put in a spreadsheet cell.

    Numbers, dates and None pass through unchanged. Text that Excel, LibreOffice
    or Google Sheets would read as a formula gets a leading apostrophe.
    """
    if value is None or isinstance(value, (bool, int, float, Decimal, date, datetime)):
        return value
    text = str(value)
    # Apostrophes already in front are quoted too, so restore() gives back exactly this text.
    if text.lstrip("'").startswith(_DANGEROUS_START):
        return "'" + text
    return text


def restore(value):
    """What neutralise() was given: its apostrophe taken off, for reading a sheet back."""
    if isinstance(value, str) and value.startswith("'") and value.lstrip("'").startswith(_DANGEROUS_START):
        return value[1:]
    return value


def xlsx_value(value):
    """
    neutralise() for an XLSX cell written with put() or append_row(). The
    cell is stored as text, which no spreadsheet runs, so the apostrophe -
    which Excel showed in the cell - is left off. Text that already starts
    with one keeps neutralise()'s, so restore() still gives back exactly what
    was written.
    """
    if isinstance(value, str) and not value.startswith("'"):
        return value
    return neutralise(value)


def _as_text(cell):
    # openpyxl takes any text starting with "=" for a formula.
    if cell.data_type == "f":
        cell.data_type = "s"
    return cell


def put(cell, value):
    """Write value into an openpyxl cell: text stays text, never a formula."""
    cell.value = xlsx_value(value)
    return _as_text(cell)


def append_row(ws, values) -> None:
    """ws.append(values), with text kept as text (see put)."""
    ws.append([xlsx_value(v) for v in values])
    for cell in ws[ws._current_row]:
        _as_text(cell)


def default_filename(name: str, extension: str = "csv") -> str:
    """'licences_2026-09-30.csv' - safe on every file system."""
    base = _SAFE_FILENAME.sub("_", str(name or "export")).strip(" ._") or "export"
    return f"{base}_{date.today().isoformat()}.{extension}"


def _csv_cell(value) -> str:
    value = neutralise(value)
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def write_csv(path, headers: Sequence[str], rows: Iterable[Sequence]) -> int:
    count = 0
    with open(path, "w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow([_csv_cell(h) for h in headers])
        for row in rows:
            writer.writerow([_csv_cell(v) for v in row])
            count += 1
    return count


def write_xlsx(path, headers: Sequence[str], rows: Iterable[Sequence], sheet: str = "Export") -> int:
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font

    def cell(value):
        if isinstance(value, Decimal):
            return float(value)
        # A control character pasted from a mail made openpyxl refuse the whole export.
        return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value

    book = Workbook()
    ws = book.active
    ws.title = (sheet or "Export")[:31]
    append_row(ws, [cell(h) for h in headers])
    for header in ws[1]:
        header.font = Font(bold=True)
    count = 0
    for row in rows:
        append_row(ws, [cell(v) for v in row])
        count += 1
    # A header row that stays put, and columns wide enough to read.
    ws.freeze_panes = "A2"
    for column in ws.columns:
        width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        ws.column_dimensions[column[0].column_letter].width = min(max(width + 2, 8), 60)
    book.save(path)
    return count


def export_rows(path, headers: Sequence[str], rows: Iterable[Sequence]) -> int:
    """Write CSV or XLSX by the file's extension. Returns rows written."""
    suffix = Path(str(path)).suffix.lower()
    rows = list(rows)
    if suffix == ".xlsx":
        return write_xlsx(path, headers, rows)
    return write_csv(path, headers, rows)


def rows_from_qtable(table) -> Tuple[List[str], List[List]]:
    """
    The headers and cell texts a QTableWidget or QTableView is showing:
    visible columns in their on-screen order, visible rows in their sorted
    order. Cell text is taken as displayed, so formats match the screen.
    """
    model = table.model()
    header = table.horizontalHeader()
    columns = []
    for visual in range(header.count()):
        logical = header.logicalIndex(visual)
        if table.isColumnHidden(logical):
            continue
        columns.append(logical)

    from PySide6.QtCore import Qt
    headers = [str(model.headerData(c, Qt.Orientation.Horizontal, Qt.ItemDataRole.DisplayRole) or "")
               for c in columns]
    rows = []
    for r in range(model.rowCount()):
        if table.isRowHidden(r):
            continue
        rows.append([model.index(r, c).data(Qt.ItemDataRole.DisplayRole) for c in columns])
    return headers, rows
