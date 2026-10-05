"""
The shared formats every screen uses: dates, money, people's names and table
export. One way each, tested once here.
"""

from datetime import date, datetime, timedelta
from decimal import Decimal

import pytest

from slate.core.domain import dates, money, table_export


@pytest.fixture(scope="module")
def qapp():
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


# ================================================================ dates


def test_one_date_format_everywhere():
    assert dates.format_date(date(2026, 10, 3)) == "3 Oct 2026"
    assert dates.format_date("2026-10-03") == "3 Oct 2026"
    assert dates.format_date(datetime(2026, 10, 3, 14, 5)) == "3 Oct 2026"
    assert dates.format_date("2026-10-03", weekday=True) == "Sat 3 Oct 2026"


def test_every_old_spelling_is_understood():
    for text in ("2026-10-03", "03-10-2026", "03/10/2026", "3 Oct 2026",
                 "Sat 3 Oct 2026", "2026-10-03 00:00:00", "3 October 2026"):
        assert dates.parse_date(text) == date(2026, 10, 3), text


def test_missing_and_unreadable_dates():
    assert dates.format_date(None) == ""
    assert dates.format_date("", empty="-") == "-"
    # Not a date: shown as it is rather than hidden.
    assert dates.format_date("sometime") == "sometime"


def test_a_timestamp_never_shows_microseconds():
    """IT-141: '2026-09-17 23:02:49.236319'."""
    assert dates.format_datetime("2026-09-17 23:02:49.236319") == "17 Sep 2026, 23:02"
    assert dates.format_datetime(datetime(2026, 9, 17, 23, 2, 49), seconds=True) == "17 Sep 2026, 23:02:49"


def test_how_long_ago():
    now = datetime(2026, 9, 30, 12, 0)
    assert dates.format_age(now - timedelta(seconds=20), now=now) == "just now"
    assert dates.format_age(now - timedelta(minutes=5), now=now) == "5 min ago"
    assert dates.format_age(now - timedelta(hours=3), now=now) == "3 h ago"
    assert dates.format_age(now - timedelta(days=1), now=now) == "yesterday"
    assert dates.format_age(now - timedelta(days=4), now=now) == "4 days ago"
    assert dates.format_age(now - timedelta(days=20), now=now) == "10 Sep 2026"


def test_date_columns_sort_by_date_not_text():
    keys = sorted(["2026-10-03", "2026-09-12", None], key=dates.date_sort_key)
    assert keys == [None, "2026-09-12", "2026-10-03"]


def test_weeks_start_on_monday():
    assert dates.WEEK_START == 0
    assert dates.week_start(date(2026, 10, 4)) == date(2026, 9, 28)  # Sunday -> Monday before
    assert dates.week_start(date(2026, 9, 28)) == date(2026, 9, 28)


def test_ranges():
    assert dates.format_range("2026-10-03", "2026-10-07") == "3 – 7 Oct 2026"
    assert dates.format_range("2026-09-28", "2026-10-02") == "28 Sep – 2 Oct 2026"
    assert dates.format_range("2026-10-03", "2026-10-03") == "3 Oct 2026"


def test_the_qt_date_edit_matches(qapp):
    from PySide6.QtWidgets import QDateEdit
    from PySide6.QtCore import QDate, Qt
    from slate.gui.core.data_display import setup_date_edit, date_item
    edit = setup_date_edit(QDateEdit(QDate(2026, 10, 3)))
    assert edit.text() == "3 Oct 2026"
    assert edit.calendarPopup()
    assert edit.calendarWidget().firstDayOfWeek() == Qt.DayOfWeek.Monday
    early, late = date_item("2026-09-12"), date_item("2026-10-03")
    assert early.text() == "12 Sep 2026" and early < late


# ================================================================ money


def test_rupees_use_indian_grouping():
    assert money.format_money(20700000, "INR") == "₹2,07,00,000.00"
    assert money.format_money(198770.6, "INR") == "₹1,98,770.60"
    assert money.format_money(999, "INR") == "₹999.00"
    assert money.format_money(-1500, "INR") == "-₹1,500.00"


def test_other_currencies_group_in_thousands():
    assert money.format_money(207000000, "USD") == "$207,000,000.00"
    assert money.format_money(1234.5, "EUR") == "€1,234.50"
    assert money.format_money(1234.5, "GBP") == "£1,234.50"


def test_compact_forms():
    assert money.format_money(20700000, "INR", compact=True) == "₹2.07 Cr"
    assert money.format_money(3850000, "INR", compact=True) == "₹38.5 L"
    assert money.format_money(45000, "INR", compact=True) == "₹45,000"
    assert money.format_money(380328958, "USD", compact=True) == "$380.3M"
    assert money.format_money(12500, "USD", compact=True) == "$12.5K"


def test_money_is_exact():
    assert money.sum_money([0.1, 0.2]) == Decimal("0.3")
    assert money.quantize("2.345") == Decimal("2.35")
    assert money.percent_of(1000, 18) == Decimal("180.00")


def test_typed_amounts_are_understood():
    assert money.parse_money("₹2,07,00,000.50") == Decimal("20700000.50")
    assert money.parse_money("$ 1,234") == Decimal("1234.00")
    assert money.parse_money("(500)") == Decimal("-500.00")
    with pytest.raises(ValueError):
        money.parse_money("lots")


def test_currencies_are_never_added_together():
    totals = money.sum_by_currency([(100, "INR"), (50, "usd"), (25, "INR"), (1, None)])
    assert totals == {"INR": Decimal(126), "USD": Decimal(50)}
    assert list(totals) == ["INR", "USD"]
    assert money.format_totals(totals) == "₹126.00 + $50.00"


def test_studio_defaults_follow_the_plan(mock_db):
    from slate.core.infra.studio_settings import StudioSettings
    StudioSettings.invalidate()
    assert money.studio_currency() == "INR"
    assert money.day_rate("INR") == Decimal(8000)
    assert money.day_rate("USD") == Decimal(300)
    assert money.day_rate("EUR") is None
    assert money.default_tax_rate("INR") == Decimal(18)
    assert money.default_tax_rate("USD") == Decimal(0)


def test_a_money_cell_sorts_by_amount(qapp):
    from slate.gui.core.data_display import money_item
    small, big = money_item(900, "INR"), money_item(10000, "INR")
    assert small < big
    assert big.text() == "₹10,000.00"


# ================================================================ people


def _people(db):
    db.execute_update("INSERT INTO ut_users (username, display_name) VALUES ('rahul.s', 'Rahul Sharma')")
    db.execute_update("INSERT INTO ut_users (username, display_name) VALUES ('arjun.p', 'Arjun P')")
    db.execute_update("INSERT INTO ut_users (username, display_name) VALUES ('admin', 'Administrator')")
    db.execute_update("INSERT INTO ut_users (username, display_name) VALUES ('nameless', '')")
    db.execute_update("INSERT INTO ut_users (username, display_name, last_day) "
                      "VALUES ('left.guy', 'Left Guy', %s)", (date.today() - timedelta(days=5),))


@pytest.fixture(params=["sqlite", "postgres"])
def db(request):
    from slate.core.domain.people import Directory
    Directory.refresh()
    if request.param == "sqlite":
        yield request.getfixturevalue("mock_db")
    else:
        yield request.getfixturevalue("pg_db")
    Directory.refresh()


def test_names_not_logins(db):
    from slate.core.domain import people
    _people(db)
    assert people.display_name("rahul.s", db) == "Rahul Sharma"
    assert people.display_name("RAHUL.S", db) == "Rahul Sharma"
    assert people.display_name("nameless", db) == "nameless"
    assert people.display_name("gone.forever", db) == "gone.forever"
    assert people.label("rahul.s", db) == "Rahul Sharma (rahul.s)"


def test_the_picker_is_alphabetical_without_leavers_or_service_accounts(db):
    """IT-020: 43 usernames in joining order, admin first, a leaver included."""
    from slate.core.domain import people
    _people(db)
    names = [p.username for p in people.people_for_picker(db)]
    assert names == ["arjun.p", "nameless", "rahul.s"]
    assert "left.guy" in [p.username for p in people.people_for_picker(db, include_leavers=True)]
    assert "admin" in [p.username for p in people.people_for_picker(db, include_service=True)]
    assert [p.username for p in people.find("sharma", db)] == ["rahul.s"]


# ================================================================ export


def test_formula_injection_is_neutralised():
    assert table_export.neutralise('=HYPERLINK("http://x")') == '\'=HYPERLINK("http://x")'
    for text in ("+cmd", "-2+3", "@SUM(A1)", "\t=1"):
        assert table_export.neutralise(text).startswith("'")
    assert table_export.neutralise(-5) == -5
    assert table_export.neutralise("Rahul") == "Rahul"
    assert table_export.neutralise(None) is None


def test_csv_opens_in_excel_with_unicode(tmp_path):
    path = tmp_path / "out.csv"
    rows = [["राहुल", "₹1,000", "=1+1"], ["emoji \U0001F3AC", 5, None]]
    assert table_export.export_rows(path, ["Name", "Amount", "Note"], rows) == 2
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    text = raw.decode("utf-8-sig")
    assert "राहुल" in text and "'=1+1" in text


def test_xlsx_keeps_numbers_as_numbers(tmp_path):
    from openpyxl import load_workbook
    path = tmp_path / "out.xlsx"
    table_export.export_rows(path, ["Name", "Seats"], [["Nuke", 12], ["=evil()", Decimal("2.5")]])
    ws = load_workbook(path).active
    assert ws["B2"].value == 12 and ws["B3"].value == 2.5
    assert (ws["A3"].value, ws["A3"].data_type) == ("=evil()", "s")


def test_export_takes_what_the_table_shows(qapp):
    from PySide6.QtWidgets import QTableWidget, QTableWidgetItem
    table = QTableWidget(3, 3)
    table.setHorizontalHeaderLabels(["ID", "Name", "Seats"])
    for r, (i, n, s) in enumerate([(1, "Nuke", "12"), (2, "Maya", "3"), (3, "Houdini", "5")]):
        for c, v in enumerate((str(i), n, s)):
            table.setItem(r, c, QTableWidgetItem(v))
    table.hideColumn(0)
    table.setRowHidden(1, True)
    headers, rows = table_export.rows_from_qtable(table)
    assert headers == ["Name", "Seats"]
    assert rows == [["Nuke", "12"], ["Houdini", "5"]]


def test_a_safe_default_filename():
    name = table_export.default_filename("IT / Licences: all")
    assert name.endswith(".csv") and "/" not in name and ":" not in name
