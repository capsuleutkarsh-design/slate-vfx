"""
Attendance, round 2 of the audit (HR2-0xx): one count for every screen, and
the corrections that lost data.

    correcting a two-session day merged the sessions                  HR2-001
    leavers absent after their last day, joiners before joining       HR2-002/003
    the hero's late count differed from the tables and the export     HR2-004
    a first-half leave day counted late                               HR2-005
    biometric times did not replace an auto punch-out                 HR2-006
    biometric import over a two-session day kept the old hours        HR2-007
    a holiday inside approved leave shown as a leave day              HR2-009
    a forgotten punch-out on a Sunday shown as worked                 HR2-010
    biometric days after a leaver's last day imported silently        HR2-012
    the export coded auto punch-outs and corrections as 'P'           HR2-014
    'Ends next day' accepted on a normal day                          HR2-023
    the export header shaded Sundays but not holidays                 HR2-025
    signing in as a service account punched it in                     HR2-018
"""

import json
from datetime import date, datetime

import pytest

from slate.core.domain import attendance_rules as rules
from slate.core.domain import biometric_import as bio
from slate.core.domain import leave_policy as lp
from slate.core.domain.central_attendance import CentralAttendance


@pytest.fixture(params=["sqlite", "postgres"])
def att(request):
    db = request.getfixturevalue("mock_db" if request.param == "sqlite" else "pg_db")
    lp.set_overrides({})
    yield CentralAttendance(db)
    lp.set_overrides({})


def _put(att, user, day, t_in, t_out=None, meta=None):
    att.db.execute_update(
        "INSERT INTO attendance_log (user_id, day_date, punch_in, punch_out, pc_name, metadata) "
        "VALUES (%s, %s, %s, %s, %s, %s)",
        (user, day.isoformat(), t_in, t_out, "TEST", json.dumps(meta or {})))


TWO = [{"in": "09:00:00", "out": "13:00:00"}, {"in": "17:00:00", "out": "19:00:00"}]


# ----------------------------------------------------------------- HR2-001

def test_correcting_a_two_session_day_keeps_the_sessions_and_the_history(att):
    day = date(2026, 9, 8)
    _put(att, "vihaan", day, "09:00:00", "19:00:00", {"sessions": TWO})
    ok, msg = att.update_record("vihaan", 2026, 9, 8, "09:00", "19:00", editor="hr.meera",
                                reason="add a reason",
                                sessions=[("09:00", "13:00"), ("17:00", "19:00")])
    assert ok, msg
    entry = att.get_user_month("vihaan", 2026, 9)["08"]
    assert rules.day_hours(entry) == pytest.approx(6.0)
    assert len(entry["sessions"]) == 2
    was = entry["edit_history"][-1]
    assert was["was_sessions"] == [["09:00", "13:00"], ["17:00", "19:00"]]

    # And a session can be removed (the break was a mistake): one session.
    ok, _ = att.update_record("vihaan", 2026, 9, 8, "09:00", "19:00", editor="hr.meera",
                              reason="one day", sessions=[("09:00", "19:00")])
    entry = att.get_user_month("vihaan", 2026, 9)["08"]
    assert ok and entry["sessions"] == [] and rules.day_hours(entry) == pytest.approx(10.0)


def test_session_rules():
    assert rules.sessions_problem([("09:00", "13:00"), ("12:00", "18:00")]).startswith(
        "Session 2 starts before session 1 ends")
    assert "only the last session" in rules.sessions_problem([("09:00", ""), ("14:00", "18:00")])
    assert rules.sessions_problem([("09:00", "13:00"), ("14:00", "")]) == ""
    # HR2-023: 'Ends next day' only when the out is not after the in.
    assert "Untick 'Ends next day'" in rules.sessions_problem([("09:00", "18:00")], overnight=True)
    assert "Untick 'Ends next day'" in rules.sessions_problem([("09:00", "09:00")], overnight=True)
    assert rules.sessions_problem([("20:00", "05:30")], overnight=True) == ""


def test_overnight_on_a_normal_day_is_refused_by_the_record_too(att):
    ok, msg = att.update_record("aarav", 2026, 9, 9, "09:00", "18:00", editor="hr", reason="x",
                                overnight=True)
    assert not ok and "Ends next day" in msg


# ------------------------------------------------------------- HR2-002/003

def test_nobody_is_absent_outside_their_employment():
    from slate.gui.attendance_export_worker import summarise
    now = datetime(2026, 9, 30, 20)
    leaver = {"holidays": set(), "expected": rules.expected_window({"last_day": "2026-09-15"})}
    joiner = {"holidays": set(), "expected": rules.expected_window({"joined_on": date(2026, 9, 14)})}
    # 1-15 Sep has 13 working days (Sundays 6 and 13 off), none punched.
    assert summarise(leaver, {}, 2026, 9, now=now)["absent"] == 13
    # 14-29 Sep: 14 working days (30 Sep is today, not yet absent).
    assert summarise(joiner, {}, 2026, 9, now=now)["absent"] == 14
    service = {"holidays": set(), "expected": rules.expected_window({}, service=True)}
    assert summarise(service, {}, 2026, 9, now=now)["absent"] == 0
    assert rules.expected_window({"deactivated_on": "2026-09-10", "last_day": "2026-09-20"})[1] \
        == date(2026, 9, 10)


# --------------------------------------------------------------- HR2-004/005

def test_late_is_one_count_for_every_screen():
    log = {"15": {"in": "11:30", "out": "19:30", "auto_logout": True}}
    summary = rules.month_summary(log, 2026, 9, now=datetime(2026, 9, 30, 20))
    day = summary["days"][14]
    assert day["state"] == rules.AUTO and day["late"] is True
    assert summary["late"] == 1


def test_a_first_half_leave_day_is_not_late():
    entry = {"in": "14:00", "out": "19:00"}
    first_half = {"type": "Sick", "half": "First half", "days": 0.5}
    second_half = {"type": "Sick", "half": "Second half", "days": 0.5}
    day = date(2026, 9, 16)
    assert not rules.arrived_late(entry, day, leave=first_half)
    assert rules.day_state(entry, day, date(2026, 9, 30), leave=first_half) != rules.LATE
    morning = {"in": "11:30", "out": "14:00"}
    assert rules.arrived_late(morning, day, leave=second_half)
    assert not rules.arrived_late(morning, day, leave={"type": "Casual", "half": "", "days": 1})


# ----------------------------------------------------------------- HR2-010

def test_a_forgotten_punch_out_on_a_sunday_is_missing_not_worked():
    sunday = date(2026, 9, 13)
    assert rules.day_state({"in": "10:00", "out": ""}, sunday, date(2026, 9, 20)) == rules.MISSING_OUT
    assert rules.day_state({"in": "10:00", "out": "16:00"}, sunday, date(2026, 9, 20)) \
        == rules.WORKED_OFF


# ------------------------------------------------------------- HR2-006/007

def _record(user, day, t_in, t_out):
    return bio.DayRecord(code=user, user_id=user, day=day, punch_in=t_in, punch_out=t_out,
                         punches=2, name="")


def test_the_machine_replaces_an_automatic_punch_out(att):
    day = date(2026, 9, 9)
    _put(att, "aarav", day, "10:00:00", "19:30:00", {"auto_logout": True, "cutoff": "19:30:00"})
    _put(att, "diya", date(2026, 9, 10), "21:00:00", None, {"missing_punch_out": True})
    result = bio.apply_days([_record("aarav", day, "10:00:00", "18:10:00"),
                             _record("diya", date(2026, 9, 10), "21:00:00", "23:40:00")], att)
    assert result["written"] == 2
    entry = att.get_user_month("aarav", 2026, 9)["09"]
    assert entry["out"] == "18:10" and not entry["auto_logout"]
    assert rules.day_state(entry, day, date(2026, 9, 30)) != rules.AUTO
    other = att.get_user_month("diya", 2026, 9)["10"]
    assert other["out"] == "23:40" and not other["missing_punch_out"]


def test_the_machine_over_a_two_session_day_keeps_times_and_hours_together(att):
    day = date(2026, 9, 8)
    _put(att, "vihaan", day, "09:00:00", "19:00:00", {"sessions": TWO})
    bio.apply_days([_record("vihaan", day, "08:30:00", "20:00:00")], att)
    entry = att.get_user_month("vihaan", 2026, 9)["08"]
    assert (entry["in"], entry["out"]) == ("08:30", "20:00")
    # 08:30-13:00 and 17:00-20:00: the break is kept, the machine's ends used.
    assert rules.day_hours(entry) == pytest.approx(7.5)


# ----------------------------------------------------------------- HR2-012

def test_days_outside_employment_are_found():
    users = {"Kabir.Left": {"joined_on": "2020-01-01", "last_day": "2026-09-15"}}
    days = [_record("kabir.left", date(2026, 9, 14), "09:00:00", "18:00:00"),
            _record("kabir.left", date(2026, 9, 25), "09:00:00", "18:00:00")]
    assert [d.day for d in bio.outside_employment(days, users)] == [date(2026, 9, 25)]


# ------------------------------------------------------------- HR2-014/025

def test_export_codes_and_header(tmp_path):
    from openpyxl import load_workbook
    from slate.gui.attendance_export_worker import build_workbook, day_code
    assert day_code(rules.AUTO, {}) == "AU"
    assert day_code(rules.PRESENT, {"corrected": True}) == "P*"
    holiday = date(2026, 9, 22)
    rows = [{"username": "aarav", "name": "Aarav", "holidays": {holiday}}]
    data = {"aarav": {"15": {"in": "10:00", "out": "19:30", "auto_logout": True}}}
    path = tmp_path / "a.xlsx"
    build_workbook(str(path), 2026, 9, rows, data, now=datetime(2026, 9, 30, 20),
                   studio_holidays={holiday})
    ws = load_workbook(path).active
    headers = [c.value for c in ws[1]]
    assert ws.cell(row=1, column=headers.index("22 Tue") + 1).fill.start_color.rgb.endswith("7F7F7F")
    assert ws.cell(row=2, column=headers.index("15 Tue") + 1).value.startswith("AU")


# ----------------------------------------------------------------- HR2-009

def test_a_holiday_inside_approved_leave_stays_a_holiday(mock_db):
    from slate.core.infra.leave_repository import LeaveRepository
    repo = LeaveRepository(mock_db)
    repo.add_holiday(date(2026, 9, 22), "Studio day", "All")
    mock_db.execute_update(
        "INSERT INTO leave_requests (user_id, start_date, end_date, type, status, half_day, "
        "days_charged) VALUES (%s, %s, %s, %s, %s, %s, %s)",
        ("vihaan", date(2026, 9, 21), date(2026, 9, 23), "Casual", lp.STATUS_APPROVED, False, 2.0))
    leave = repo.approved_leave(date(2026, 9, 1), date(2026, 9, 30), ["vihaan"])["vihaan"]
    assert sorted(leave) == [date(2026, 9, 21), date(2026, 9, 23)]


# ----------------------------------------------------------------- HR2-018

def test_signing_in_never_punches_in_a_service_account(att):
    from slate.core.domain import people
    people.refresh()
    assert att.log_action("admin", "in", automatic=True) is None
    assert att.today_state("admin")["state"] == "out"
    people.refresh()
