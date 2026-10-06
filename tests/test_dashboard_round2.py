"""
Round-2 dashboard fixes (DSH2-*): the one save path, artists' own statuses,
projects, versions and the review queue, the production summary, the Excel
backup and the screen - on an isolated SQLite database.
"""

import json
from datetime import date

import pytest
from PySide6.QtCore import QSettings

from slate.gui.tabs.vfx_dashboard_pro.core.sqlite_handler import (
    ProjectClosedError, SQLiteHandler, StaleDataError,
)
from slate.gui.tabs.vfx_dashboard_pro.models.shot_model import Shot
from tests.dashboard_util import open_project

PROJECT = "R2P"


@pytest.fixture(autouse=True)
def online(monkeypatch):
    import slate.core.domain.access as access
    monkeypatch.setattr(access, "is_offline_fallback", lambda: False)


@pytest.fixture(autouse=True)
def scratch_settings(tmp_path, monkeypatch):
    from slate.gui.tabs.vfx_dashboard_pro.ui.components import column_layout_manager as clm
    path = str(tmp_path / "layouts.ini")
    monkeypatch.setattr(clm, "settings_factory", lambda: QSettings(path, QSettings.Format.IniFormat))


def _shot(name="SH010", reel="R01", **kw):
    shot = Shot(shot_name=name, reel_episode=reel, status="WIP", **kw)
    shot.dept("comp").artist = "priya"
    shot.dept("comp").status = "WIP"
    return shot


def _handler(db, role="supervisor", **kw):
    kw.setdefault("username", "sup")
    return SQLiteHandler(PROJECT, db_manager=db, user_role=role, **kw)


@pytest.fixture
def seeded(mock_db):
    open_project(mock_db, PROJECT)
    handler = _handler(mock_db)
    assert handler.write_shots([_shot("SH010"), _shot("SH020")])
    return handler


def _history(db, field=None):
    rows = db.execute_query("SELECT * FROM change_history WHERE project_code=%s ORDER BY id",
                            (PROJECT,), fetch="all") or []
    rows = [dict(r) for r in rows]
    return [r for r in rows if field is None or r["field_changed"] == field]


def _task(db, shot_id, dept):
    return dict(db.execute_query("SELECT * FROM tracking_tasks WHERE shot_id=%s AND department=%s",
                                 (shot_id, dept), fetch="one") or {})


# ------------------------------------------------------------------ the save path
class TestSavePath:

    def test_every_saved_field_is_in_the_history(self, seeded, mock_db):
        """DSH2-005 (regression of DSH-005): version, frames, bid and actual days are logged."""
        shot = seeded.read_shots()[0]
        shot.curr_version = "v500"
        shot.edit_frames = 96.0
        shot.dept("roto").bid_days = 2.5
        shot.dept("comp").actual_days = 1.25
        assert seeded.write_shots([shot])
        fields = {r["field_changed"]: r for r in _history(mock_db)}
        assert fields["curr_version"]["new_value"] == "v500"
        assert fields["edit_frames"]["new_value"] == "96"
        assert fields["roto_bid_days"]["new_value"] == "2.5"
        assert fields["comp_actual_days"]["new_value"] == "1.25"
        assert fields["curr_version"]["user_id"] == "sup"

    def test_only_changed_department_rows_are_written_and_actual_days_stay_unrecorded(self, seeded, mock_db):
        """DSH2-008 / DSH2-054."""
        shot = seeded.read_shots()[0]
        mock_db.execute_update("UPDATE tracking_tasks SET bid_days = 7 WHERE shot_id=%s AND department='roto'",
                               (shot.id,))
        shot = seeded.read_shots()[0]
        shot.sow = "only the scope"
        assert seeded.write_shots([shot])
        assert _task(mock_db, shot.id, "roto")["bid_days"] == 7
        assert _task(mock_db, shot.id, "comp").get("actual_days") is None
        again = seeded.read_shots()[0]
        assert again.dept("comp").actual_days is None

    def test_a_cleared_artist_stays_cleared(self, seeded):
        """DSH2-007: it was refilled from Comp on every read."""
        shot = seeded.read_shots()[0]
        shot.assigned_artist = "someone"
        seeded.write_shots([shot])
        shot = seeded.read_shots()[0]
        shot.assigned_artist = ""
        seeded.write_shots([shot])
        assert seeded.read_shots()[0].assigned_artist == ""

    def test_an_eta_never_becomes_the_target(self, seeded):
        """DSH2-048."""
        shot = seeded.read_shots()[0]
        shot.dept("prep").eta = "2026-11-30"
        seeded.write_shots([shot])
        after = seeded.read_shots()[0]
        assert after.dept("prep").target in ("", None)
        assert after.dept("prep").eta == "2026-11-30"

    def test_a_save_into_a_deleted_project_is_refused_and_brings_nothing_back(self, seeded, mock_db):
        """DSH2-014."""
        shot = seeded.read_shots()[0]
        from slate.gui.tabs.vfx_dashboard_pro.core.project_manager import ProjectManager
        manager = ProjectManager()
        assert manager.delete_project(PROJECT)
        shot.sow = "after the delete"
        with pytest.raises(ProjectClosedError):
            seeded.write_shots([shot])
        assert mock_db.execute_query("SELECT * FROM tracking_projects WHERE code=%s", (PROJECT,),
                                     fetch="all") == []
        assert mock_db.execute_query("SELECT * FROM tracking_shots WHERE project_code=%s", (PROJECT,),
                                     fetch="all") == []

    def test_a_save_into_an_archived_project_is_refused(self, seeded, mock_db):
        """DSH2-066."""
        shot = seeded.read_shots()[0]
        mock_db.execute_update("UPDATE tracking_projects SET active = 0 WHERE code=%s", (PROJECT,))
        shot.sow = "x"
        with pytest.raises(ProjectClosedError) as caught:
            seeded.write_shots([shot])
        assert caught.value.archived

    def test_a_shot_deleted_elsewhere_is_a_conflict_not_reinserted(self, seeded, mock_db):
        """DSH2-014: a pending shot whose row vanished."""
        shot = seeded.read_shots()[0]
        mock_db.execute_update("DELETE FROM tracking_shots WHERE id=%s", (shot.id,))
        shot.sow = "x"
        with pytest.raises(StaleDataError) as caught:
            seeded.write_shots([shot])
        assert caught.value.conflicts[0]["kind"] == "deleted"
        assert len(seeded.read_shots()) == 1

    def test_two_people_adding_the_same_shot_is_a_conflict(self, seeded):
        """DSH2-067."""
        mine = _shot("SH999")
        theirs = _shot("SH999")
        theirs.sow = "theirs"
        assert seeded.write_shots([theirs])
        mine.sow = "mine"
        with pytest.raises(StaleDataError) as caught:
            seeded.write_shots([mine, _shot("SH998")])
        assert caught.value.conflicts[0]["kind"] == "added"
        names = {s.shot_name: s for s in seeded.read_shots()}
        assert names["SH999"].sow == "theirs" and "SH998" not in names

    def test_a_batch_save_keeps_its_version_lock(self, seeded, mock_db):
        """DSH2-057: a save made in between is a conflict, and nothing of the batch is written."""
        shots = seeded.read_shots()
        other = _handler(mock_db).read_shots()
        other[0].sow = "in between"
        assert _handler(mock_db).write_shots([other[0]])
        for s in shots:
            s.sow = "batch"
        with pytest.raises(StaleDataError):
            seeded.write_shots(shots)
        after = {s.shot_name: s.sow for s in seeded.read_shots()}
        assert after[other[0].shot_name] == "in between"
        assert "batch" not in after.values()

    def test_a_failed_read_raises(self, mock_db, monkeypatch):
        """DSH2-016: a read failure must not look like an empty project."""
        open_project(mock_db, PROJECT)
        handler = _handler(mock_db)
        monkeypatch.setattr(mock_db, "execute_query", lambda *a, **k: None)
        with pytest.raises(RuntimeError):
            handler.read_shots()

    def test_a_lead_is_not_blamed_for_someone_elses_change(self, seeded, mock_db):
        """DSH2-012: a stale lead save is a conflict, not a permission breach."""
        from tests.dashboard_util import person
        lead = _handler(mock_db, role=["lead"], username=person(mock_db, "rl", "Roto Lead"))
        mine = lead.read_shots()[0]
        from slate.gui.tabs.vfx_dashboard_pro.ui.shot_table_model import snapshot
        mine._baseline = snapshot(mine)
        other = seeded.read_shots()[0]
        other.priority = 0
        assert seeded.write_shots([other])
        mine.dept("roto").status = "WIP"
        with pytest.raises(StaleDataError):
            lead.write_shots([mine])


# ------------------------------------------------------------------ artists
class TestArtistStatus:

    def _artist(self, db):
        return _handler(db, role="artist", username="priya")

    def test_a_status_change_keeps_bid_days_and_targets(self, seeded, mock_db):
        """DSH2-013."""
        shot = seeded.read_shots()[0]
        mock_db.execute_update(
            "UPDATE tracking_tasks SET bid_days = 5, target_date = '2026-12-01' WHERE shot_id=%s "
            "AND department='comp'", (shot.id,))
        assert self._artist(mock_db).update_department_status(
            "SH010", "R01", "comp", "SENT FOR REVIEW", 0, actor_identities=["priya"], shot_id=shot.id)
        task = _task(mock_db, shot.id, "comp")
        assert task["status"] == "SENT FOR REVIEW"
        assert task["bid_days"] == 5 and task["target_date"] == "2026-12-01"
        assert _history(mock_db, "comp_status")[-1]["new_value"] == "SENT FOR REVIEW"

    def test_an_undo_may_put_back_a_status_the_artist_cannot_choose(self, seeded, mock_db):
        """DSH2-003: N/A -> WIP by mistake, and back."""
        shot = seeded.read_shots()[0]
        shot.dept("comp").status = "N/A"
        seeded.write_shots([shot])
        artist = self._artist(mock_db)
        assert artist.update_department_status("SH010", "R01", "comp", "WIP", 0,
                                               actor_identities=["priya"], shot_id=shot.id)
        with pytest.raises(PermissionError, match="supervisor or coordinator"):
            artist.update_department_status("SH010", "R01", "comp", "N/A", 0,
                                            actor_identities=["priya"], shot_id=shot.id)
        assert artist.update_department_status("SH010", "R01", "comp", "N/A", 0, actor_identities=["priya"],
                                               shot_id=shot.id, undo_of="WIP")
        assert seeded.read_shots()[0].dept("comp").status == "N/A"

    def test_an_undo_is_not_a_way_round_the_rule(self, seeded, mock_db):
        """Somebody else's last change cannot be 'undone' into a verdict."""
        shot = seeded.read_shots()[0]
        artist = self._artist(mock_db)
        with pytest.raises(PermissionError):
            artist.update_department_status("SH010", "R01", "comp", "APPROVED", 0, actor_identities=["priya"],
                                            shot_id=shot.id, undo_of="WIP")


# ------------------------------------------------------------------ projects
class TestProjects:

    def test_an_archived_code_is_refused_case_insensitively(self, seeded, mock_db):
        """DSH2-009 / DSH2-075."""
        from slate.gui.tabs.vfx_dashboard_pro.core.project_manager import ProjectManager
        manager = ProjectManager()
        assert manager.archive_project(PROJECT)
        assert manager.add_project(PROJECT.lower(), "New", "", "") is None
        assert "archived" in manager.last_error
        row = dict(mock_db.execute_query("SELECT active FROM tracking_projects WHERE code=%s",
                                         (PROJECT,), fetch="one"))
        assert str(row["active"]) in ("0", "False", "false")

    @pytest.mark.parametrize("code", ["KLC 2", "KLC/A", "a..b", ""])
    def test_a_bad_code_is_refused(self, mock_db, code):
        from slate.gui.tabs.vfx_dashboard_pro.core.project_manager import ProjectManager
        assert ProjectManager().add_project(code, "X", "", "") is None

    def test_a_code_is_stored_in_capitals(self, mock_db):
        from slate.gui.tabs.vfx_dashboard_pro.core.project_manager import ProjectManager
        assert ProjectManager().add_project("klc", "Kaalchakra", "", "").code == "KLC"

    def test_restoring_an_active_project_is_refused(self, seeded, mock_db):
        """DSH2-041."""
        from slate.gui.tabs.vfx_dashboard_pro.core.project_manager import ProjectManager
        manager = ProjectManager()
        assert manager.restore_project(PROJECT) is False
        assert _history(mock_db, "project") == []

    def test_a_permanent_delete_takes_versions_and_deliveries(self, seeded, mock_db):
        """DSH2-103."""
        from slate.core.domain.deliveries import DeliveryStore
        from slate.core.domain.versions import VersionStore
        from slate.gui.tabs.vfx_dashboard_pro.core.project_manager import ProjectManager
        version = VersionStore(db=mock_db).add_version(PROJECT, "SH010", reel="R01")
        VersionStore(db=mock_db).add_note(version.id, "note")
        DeliveryStore(db=mock_db).create_delivery(PROJECT, "DEL", [version.id])
        assert ProjectManager().delete_project(PROJECT)
        for table in ("tracking_versions", "tracking_deliveries"):
            assert mock_db.execute_query(f"SELECT * FROM {table} WHERE project_code=%s", (PROJECT,),
                                         fetch="all") == []
        assert mock_db.execute_query("SELECT * FROM tracking_version_notes", fetch="all") == []


# ------------------------------------------------------------------ versions
class TestVersions:

    def test_versions_belong_to_one_reel(self, mock_db):
        """DSH2-021."""
        from slate.core.domain.versions import VersionStore
        store = VersionStore(db=mock_db)
        assert store.add_version(PROJECT, "SH010", reel="R01").version_name == "v001"
        assert store.add_version(PROJECT, "SH010", reel="R02").version_name == "v001"
        assert [v.reel for v in store.list_for_shot(PROJECT, "SH010", reel="R02")] == ["R02"]

    def test_a_lead_judges_only_their_department(self, mock_db):
        """DSH2-022."""
        from slate.core.domain.versions import STATUS_APPROVED, VersionStore
        plain = VersionStore(db=mock_db)
        comp = plain.add_version(PROJECT, "SH010", department="comp")
        roto = plain.add_version(PROJECT, "SH020", department="roto")
        lead = VersionStore(db=mock_db, roles=["lead"], departments={"roto"})
        with pytest.raises(PermissionError):
            lead.update_version(comp.id, status=STATUS_APPROVED)
        assert lead.update_version(roto.id, status=STATUS_APPROVED)

    def test_the_screen_reads_a_leads_department_where_saving_does(self):
        """B3: only the Department box (job_title) on the person's own record."""
        from types import SimpleNamespace
        from slate.gui.tabs.vfx_dashboard_pro.ui.dashboard_widget import DashboardWidget
        detect = DashboardWidget._detect_user_department_family
        assert detect(SimpleNamespace(user_data={"job_title": "Producer", "department": "Roto"})) is None
        assert detect(SimpleNamespace(user_data={"job_title": "Roto Lead"})) == "roto"

    def test_a_new_version_cannot_be_born_approved_without_the_right(self, mock_db):
        """DSH2-068."""
        from slate.core.domain.versions import STATUS_APPROVED, VersionStore
        with pytest.raises(PermissionError):
            VersionStore(db=mock_db, roles=["artist"]).add_version(PROJECT, "SH010", status=STATUS_APPROVED)

    def test_the_queue_holds_sent_versions_of_live_shots(self, seeded, mock_db):
        """DSH2-031 / DSH2-098 / DSH2-035."""
        from slate.core.domain.versions import SENT_CLIENT, VersionStore
        store = VersionStore(db=mock_db)
        unsent = store.add_version(PROJECT, "SH010", reel="R01")
        sent = store.add_version(PROJECT, "SH010", reel="R01", sent_to=SENT_CLIENT)
        omitted = store.add_version(PROJECT, "SH020", reel="R01", sent_to=SENT_CLIENT)
        shot = next(s for s in seeded.read_shots() if s.shot_name == "SH020")
        shot.status = "OMIT"
        seeded.write_shots([shot])
        queue = store.awaiting_review(PROJECT)
        assert [v.id for v in queue] == [sent.id]
        assert unsent.id not in {v.id for v in queue} and omitted.id not in {v.id for v in queue}
        assert store.awaiting_review(PROJECT, visible={("r01", "sh020")}) == []


# ------------------------------------------------------------------ summary
class TestSummary:

    def test_rules(self):
        """DSH2-028 / 077 / 078 / 079."""
        from slate.core.domain.production_summary import build_summary
        done = Shot(shot_name="SH020", reel_episode="R01", status="APPROVED")
        done.dept("comp").artist = "Priya"
        done.dept("comp").status = "WIP"
        done.dept("comp").bid_days = 5
        busy = Shot(shot_name="SH030", reel_episode="R01", status="WIP", target="2026-10-01 00:00:00")
        busy.dept("comp").artist = "priya "
        busy.dept("comp").status = "WIP"
        busy.dept("comp").bid_days = 2
        lonely = Shot(shot_name="SH040", reel_episode="R02", status="WIP")
        finished_alone = Shot(shot_name="SH050", reel_episode="R02", status="APPROVED")
        summary = build_summary([done, busy, lonely, finished_alone], today=date(2026, 10, 3))
        assert summary.outstanding_bid_days == 2
        assert [(a.shots, a.outstanding_days) for a in summary.artists] == [(2, 2)]
        assert summary.unassigned == ["R02 / SH040"]
        assert [e.shot_name for e in summary.late] == ["SH030"]


# ------------------------------------------------------------------ Excel backup
class TestExcelBackup:

    def _project(self, tmp_path):
        from slate.gui.tabs.vfx_dashboard_pro.core.project_manager import ProjectConfig, default_column_mapping
        return ProjectConfig(code="XL", name="Excel", excel_path=str(tmp_path / "XL.xlsx"),
                             column_mapping=default_column_mapping())

    def _write(self, tmp_path, shots):
        from slate.gui.tabs.vfx_dashboard_pro.core.excel_handler import ExcelHandler
        project = self._project(tmp_path)
        handler = ExcelHandler(project.excel_path, project)
        assert handler.create_workbook()
        assert handler.write_shots(shots)
        return project, ExcelHandler(project.excel_path, project)

    def test_targets_actual_days_and_reels_survive_a_round_trip(self, tmp_path):
        """DSH2-023 / 070 / 071 / 072."""
        a = Shot(shot_name="SH010", reel_episode="R01", status="", target="2026-10-31", in_os="OS",
                 edit_status="Locked", description="Hero", is_hero=True, priority=1)
        a.dept("comp").actual_days = 4.0
        from slate.gui.tabs.vfx_dashboard_pro.models.shot_model import FeedbackEntry
        a.feedback_client = [FeedbackEntry(text="first note"), FeedbackEntry(text="second note")]
        b = Shot(shot_name="SH010", reel_episode="R02", status="WIP")
        _project, reader = self._write(tmp_path, [a, b])
        got = {s.reel_episode: s for s in reader.read_shots()}
        r01 = got["R01"]
        assert r01.target == "2026-10-31" and r01.in_os == "OS" and r01.edit_status == "Locked"
        assert r01.description == "Hero" and r01.status == "" and r01.priority == 1
        assert r01.dept("comp").actual_days == 4.0
        assert r01.is_hero and not got["R02"].is_hero
        assert [e.text for e in r01.feedback_client] == ["first note", "second note"]

    def test_pasted_control_characters_and_merged_cells_do_not_stop_the_backup(self, tmp_path):
        """DSH2-025 / DSH2-069."""
        from openpyxl import load_workbook
        project = self._project(tmp_path)
        from slate.gui.tabs.vfx_dashboard_pro.core.excel_handler import ExcelHandler
        handler = ExcelHandler(project.excel_path, project)
        handler.create_workbook()
        wb = load_workbook(project.excel_path)
        ws = wb.active
        sow = project.column_mapping["sow"]
        ws.merge_cells(f"{sow}3:{sow}4")
        wb.save(project.excel_path)
        shots = [Shot(shot_name="SH010", reel_episode="R01", sow="Paste from mail\x0bline two"),
                 Shot(shot_name="SH020", reel_episode="R01")]
        assert ExcelHandler(project.excel_path, project).write_shots(shots)

    def test_the_sheet_reads_as_a_passbook(self, tmp_path):
        """DSH2-046: priority names, real dates, helper sheet hidden, heading frozen."""
        from openpyxl import load_workbook
        project, _reader = self._write(tmp_path, [Shot(shot_name="SH010", reel_episode="R01", priority=0,
                                                       target="2026-10-31")])
        wb = load_workbook(project.excel_path)
        ws = wb[project.sheet_name]
        row = 3
        assert ws[f"{project.column_mapping['priority']}{row}"].value == "Urgent"
        assert ws[f"{project.column_mapping['target']}{row}"].is_date
        assert wb["Slate_DATA"].sheet_state == "hidden"
        assert ws.freeze_panes

    def test_a_failure_that_raises_is_recorded(self, tmp_path, monkeypatch):
        """DSH2-026."""
        from slate.gui.tabs.vfx_dashboard_pro.core import excel_handler
        from slate.gui.tabs.vfx_dashboard_pro.ui.dashboard_sync_service import DashboardSyncService
        project = self._project(tmp_path)

        class Manager:
            def ensure_excel_path(self, code):
                return project.excel_path
        monkeypatch.setattr(excel_handler.ExcelHandler, "write_shots",
                            lambda *a, **k: (_ for _ in ()).throw(ValueError("boom")))
        service = DashboardSyncService(Manager())
        ok, _ = service.mirror_shots_to_excel([Shot(shot_name="SH010")], project, None, force=True)
        assert not ok and service.last_backup_error == "boom"


# ------------------------------------------------------------------ the grid model
class TestGridModel:

    def test_redo_and_a_toast_only_undoes_its_own_step(self):
        """DSH2-062 / DSH2-015."""
        from slate.gui.tabs.vfx_dashboard_pro.ui.shot_table_model import ShotTableModel
        shot = Shot(shot_name="SH010", reel_episode="R01", status="WIP")
        model = ShotTableModel([shot], user_role="supervisor")
        model.apply_edit([shot], lambda s: setattr(s, "priority", 0), "priority change")
        first = model.last_step()
        model.apply_edit([shot], lambda s: setattr(s, "sow", "later"), "scope change")
        assert model.undo(first) is None                 # not the newest any more
        assert shot.priority == 0 and shot.sow == "later"
        model.undo()
        assert shot.sow == ""
        model.redo()
        assert shot.sow == "later"

    def test_an_unknown_priority_is_kept_by_the_editor(self, qtbot):
        """DSH2-055."""
        from PySide6.QtCore import Qt
        from slate.gui.tabs.vfx_dashboard_pro.ui.cell_delegates import PriorityDelegate
        from slate.gui.tabs.vfx_dashboard_pro.ui.shot_table_model import ShotTableModel
        shot = Shot(shot_name="SH010", reel_episode="R01", priority=5)
        model = ShotTableModel([shot], user_role="supervisor")
        index = model.index(0, model.column_index("priority"))
        delegate = PriorityDelegate()
        editor = delegate.createEditor(None, None, index)
        delegate.setEditorData(editor, index)
        delegate.setModelData(editor, model, index)
        assert shot.priority == 5 and not shot._modified


# ------------------------------------------------------------------ departments
@pytest.mark.parametrize("title, family", [
    ("Paint Lead", "prep"), ("Head of Paint", "prep"), ("Senior Painter", "prep"),
    ("Trainee", None), ("Roto Lead", "roto"), ("Mograph Lead", "mgfx"), ("Matte Painter", "dmp"),
])
def test_department_from_job_title(title, family):
    """DSH2-006: whole words, not pieces of words ('ai' in 'paint')."""
    from types import SimpleNamespace
    from slate.gui.tabs.vfx_dashboard_pro.ui.dashboard_widget import DashboardWidget
    fake = SimpleNamespace(user_data={"job_title": title})
    assert DashboardWidget._detect_user_department_family(fake) == family


# ------------------------------------------------------------------ filters, groups, dialogs
class TestFiltersAndGroups:

    def test_rules_match_what_the_grid_shows(self):
        """DSH2-049 / DSH2-121."""
        from slate.gui.tabs.vfx_dashboard_pro.controllers.filter_mixin import rule_passes
        shot = Shot(shot_name="SH010", target="2026-12-24", priority=3)
        assert rule_passes(shot, {"field": "Target", "operator": "is", "value": "24 Dec 2026"})
        assert rule_passes(shot, {"field": "Target", "operator": "contains", "value": "Dec"})
        assert rule_passes(shot, {"field": "Priority", "operator": "Equals", "value": "3"})   # an old rule
        assert rule_passes(shot, {"field": "Priority", "operator": "is", "value": "low"})

    def test_target_sorts_dates_then_text_then_blanks(self):
        """DSH2-122."""
        from PySide6.QtCore import Qt
        from slate.gui.tabs.vfx_dashboard_pro.ui.shot_proxy_models import ShotFilterProxy
        from slate.gui.tabs.vfx_dashboard_pro.ui.shot_table_model import ShotTableModel
        shots = [Shot(shot_name=n, target=t) for n, t in
                 (("A", ""), ("B", "TBD"), ("C", "2026-01-02"), ("D", "Hold"), ("E", ""))]
        model = ShotTableModel(shots, user_role="supervisor")
        proxy = ShotFilterProxy()
        proxy.setSourceModel(model)
        proxy.sort(model.column_index("target"), Qt.SortOrder.AscendingOrder)
        assert [s.shot_name for s in proxy.shots()] == ["C", "D", "B", "A", "E"]

    def test_group_headings_follow_an_edit(self):
        """DSH2-056 / DSH2-102."""
        from slate.gui.tabs.vfx_dashboard_pro.ui.shot_proxy_models import ShotFilterProxy, ShotGroupModel
        from slate.gui.tabs.vfx_dashboard_pro.ui.shot_table_model import ShotTableModel
        a = Shot(shot_name="A", status="WIP", edit_frames=10)
        b = Shot(shot_name="B", status="OMIT", edit_frames=99)
        b.dept("comp").bid_days = 5
        model = ShotTableModel([a, b], user_role="supervisor")
        proxy = ShotFilterProxy()
        proxy.setSourceModel(model)
        groups = ShotGroupModel()
        groups.setSourceModel(proxy)
        groups.set_group_by("Status")
        model.apply_edit([a], lambda s: setattr(s, "status", "APPROVED"), "status change")
        titles = {p["title"]: p for k, p in groups._rows if k == "header"}
        assert "APPROVED" in titles and "WIP" not in titles
        groups.set_group_by("Reel / Sequence")
        header = next(p for k, p in groups._rows if k == "header")
        assert header["total_frames"] == 10 and header["total_bids"] == 0

    def test_batch_edit_needs_a_ticked_field(self, qtbot):
        """DSH2-044."""
        from slate.gui.tabs.vfx_dashboard_pro.ui.batch_edit_dialog import BatchEditDialog
        dialog = BatchEditDialog(3)
        qtbot.addWidget(dialog)
        assert not dialog.apply_btn.isEnabled()
        dialog.target_cb.setChecked(True)
        assert dialog.apply_btn.isEnabled()
        assert dialog.get_updates() == {"target": ""}

    def test_add_shots_lists_what_it_skips(self, qtbot):
        """DSH2-047."""
        from slate.gui.tabs.vfx_dashboard_pro.ui.add_shots_dialog import AddShotsDialog
        dialog = AddShotsDialog(existing_reels=["R01"], existing_shots=[("R01", "SH010")])
        qtbot.addWidget(dialog)
        dialog.reel_input.setCurrentText("R01")
        dialog.shots_input.setPlainText("SH010\nNEW_A\nnew_a\nNEW_B")
        values = dialog.get_values()
        assert values["shots"] == ["NEW_A", "NEW_B"]
        assert values["skipped"] == ["SH010", "new_a"]
        assert "SH010" in dialog.preview_label.text()
