"""
The VFX Dashboard audit fixes (DSH-*), end to end on an isolated SQLite
database: the grid model, the save model, the board, the detail panel,
filters, permissions and the project actions.
"""

import json

import pytest
from PySide6.QtCore import QSettings, Qt

from slate.gui.tabs.vfx_dashboard_pro.core.sqlite_handler import SQLiteHandler, StaleDataError
from slate.gui.tabs.vfx_dashboard_pro.models.shot_model import Shot
from slate.gui.tabs.vfx_dashboard_pro.ui.shot_table_model import ShotTableModel

PROJECT = "DSH"


@pytest.fixture(autouse=True)
def online(monkeypatch):
    import slate.core.domain.access as access
    monkeypatch.setattr(access, "is_offline_fallback", lambda: False)


@pytest.fixture(autouse=True)
def scratch_settings(tmp_path, monkeypatch):
    """Layouts and the last project go to a scratch file, never the registry."""
    from slate.gui.tabs.vfx_dashboard_pro.ui.components import column_layout_manager as clm
    path = str(tmp_path / "layouts.ini")
    monkeypatch.setattr(clm, "settings_factory", lambda: QSettings(path, QSettings.Format.IniFormat))


def _shot(name="SH010", reel="R01", status="WIP", **kw):
    shot = Shot(shot_name=name, reel_episode=reel, status=status, **kw)
    shot.dept("comp").artist = kw.pop("comp_artist", "") if False else shot.dept("comp").artist
    return shot


def _project(mock_db, shots):
    mock_db.save_tracking_project(PROJECT, "Dashboard test", json.dumps({"code": PROJECT, "name": "Dashboard test"}))
    handler = SQLiteHandler(PROJECT, db_manager=mock_db, user_role="supervisor", username="sup")
    assert handler.write_shots(shots) is True
    return handler


def _widget(qtbot, role="Supervisor", **extra):
    from slate.gui.tabs.vfx_dashboard_pro.ui.dashboard_widget import DashboardWidget

    class Users:
        def get_all_users(self):
            return {
                "rahul": {"display_name": "Rahul", "roles": ["Artist"]},
                "priya": {"display_name": "Priya Sharma", "roles": ["Compositor"]},
                "sup": {"display_name": "Sanjay", "roles": ["Supervisor"]},
            }

    data = {"username": extra.pop("username", "sup"), "display_name": extra.pop("display_name", "Sanjay"),
            "roles": [role], "inherit_app_theme": True}
    data.update(extra)
    widget = DashboardWidget(user_data=data, user_manager=Users())
    qtbot.addWidget(widget)
    return widget


def _open(widget):
    widget.project_manager.load_config()
    widget.load_projects()
    index = widget.project_combo.findData(PROJECT)
    assert index > 0
    widget.project_combo.setCurrentIndex(index)
    return widget


def _col(model, key):
    return model.column_index(key)


# ------------------------------------------------------------------ grid model
class TestGridModel:

    def _model(self, shots=None, role="supervisor"):
        shots = shots or [_shot(sow="डिजिटल पेंटिंग – आकाश बदलें", target="2026-10-03")]
        return ShotTableModel(shots, user_role=role), shots

    def test_editors_open_on_the_stored_value_and_closing_changes_nothing(self, qtbot):
        """DSH-001: an editor opened and closed used to blank the cell."""
        from PySide6.QtWidgets import QTableView
        from slate.gui.tabs.vfx_dashboard_pro.ui.cell_delegates import TextDelegate, TargetDateDelegate
        model, shots = self._model()
        view = QTableView()
        qtbot.addWidget(view)
        view.setModel(model)
        sow = model.index(0, _col(model, "sow"))
        assert sow.data(Qt.ItemDataRole.EditRole) == shots[0].sow
        delegate = TextDelegate()
        editor = delegate.createEditor(view, None, sow)
        delegate.setEditorData(editor, sow)
        delegate.setModelData(editor, model, sow)
        target = model.index(0, _col(model, "target"))
        tdel = TargetDateDelegate()
        teditor = tdel.createEditor(view, None, target)
        tdel.setEditorData(teditor, target)
        tdel.setModelData(teditor, model, target)
        assert shots[0].sow == "डिजिटल पेंटिंग – आकाश बदलें"
        assert shots[0].target == "2026-10-03"
        assert shots[0]._modified is False

    def test_a_blank_department_status_is_not_turned_into_wip(self, qtbot):
        """Extra (triage of DSH-001): the status editor opened on WIP for a blank cell."""
        from slate.gui.tabs.vfx_dashboard_pro.ui.status_delegate import StatusDelegate
        model, shots = self._model()
        index = model.index(0, _col(model, "roto"))
        delegate = StatusDelegate()
        editor = delegate.createEditor(None, None, index)
        delegate.setEditorData(editor, index)
        assert editor.currentData() == ""
        delegate.setModelData(editor, model, index)
        assert shots[0].dept("roto").status == ""
        assert shots[0]._modified is False

    def test_six_edits_and_six_undos_leave_nothing_to_save(self):
        """DSH-017."""
        model, shots = self._model()
        for value in ("A", "B", "C"):
            model.setData(model.index(0, _col(model, "sow")), value)
        for value in ("APPROVED", "RETAKE", "YTS"):
            model.setData(model.index(0, _col(model, "status")), value)
        assert shots[0]._modified is True
        for _ in range(6):
            model.undo()
        assert shots[0]._modified is False
        assert model.pending_shots() == []

    def test_bad_numbers_are_refused_not_turned_into_zero(self):
        """DSH-044."""
        model, shots = self._model([_shot(priority=2, edit_frames=10)])
        prio, frames = model.index(0, _col(model, "priority")), model.index(0, _col(model, "frames"))
        assert model.setData(prio, "bogus") is False and shots[0].priority == 2
        assert model.setData(prio, "High") is True and shots[0].priority == 1
        assert model.setData(frames, "-5") is False
        assert model.setData(frames, "1,200") is True and shots[0].edit_frames == 1200

    def test_target_is_a_date_sorts_by_date_and_flags_overdue(self):
        """DSH-043."""
        from slate.gui.tabs.vfx_dashboard_pro.ui.shot_proxy_models import ShotFilterProxy
        shots = [_shot("A", target="2027-01-01"), _shot("B", target="TBD"),
                 _shot("C", target="2026-01-02"), _shot("D", target="")]
        model = ShotTableModel(shots, user_role="supervisor")
        assert model.setData(model.index(0, _col(model, "target")), "next friday") is False
        proxy = ShotFilterProxy()
        proxy.setSourceModel(model)
        proxy.sort(_col(model, "target"), Qt.SortOrder.AscendingOrder)
        assert [s.shot_name for s in proxy.shots()][:2] == ["C", "A"]
        proxy.sort(_col(model, "target"), Qt.SortOrder.DescendingOrder)
        assert [s.shot_name for s in proxy.shots()][:2] == ["A", "C"]
        overdue = model.index(2, _col(model, "target")).data(Qt.ItemDataRole.ForegroundRole)
        assert overdue is not None

    def test_display_is_full_and_labelled(self):
        """DSH-069/070/071/072/073/077."""
        long_sow = "x" * 120
        model, _ = self._model([_shot(sow=long_sow, priority=0)])
        assert model.index(0, _col(model, "sow")).data() == long_sow
        assert model.index(0, _col(model, "roto")).data() == ""
        assert model.index(0, _col(model, "priority")).data() == "Urgent"
        centre = int(Qt.AlignmentFlag.AlignCenter)
        assert model.index(0, _col(model, "priority")).data(Qt.ItemDataRole.TextAlignmentRole) == centre
        assert model.index(0, _col(model, "type")).data(Qt.ItemDataRole.TextAlignmentRole) != centre
        tip = model.headerData(_col(model, "matchmove"), Qt.Orientation.Horizontal, Qt.ItemDataRole.ToolTipRole)
        assert "Matchmove" in tip

    def test_edited_rows_are_marked(self):
        """DSH-075."""
        from slate.gui.tabs.vfx_dashboard_pro.ui.shot_table_model import MODIFIED_ROLE
        model, _ = self._model()
        index = model.index(0, _col(model, "sow"))
        model.setData(index, "changed")
        assert index.data(MODIFIED_ROLE) is True
        assert index.data(Qt.ItemDataRole.BackgroundRole) is not None
        assert model.index(0, _col(model, "type")).data(MODIFIED_ROLE) is False

    def test_an_artist_sees_which_cell_is_theirs(self):
        """DSH-130."""
        shot = _shot()
        shot.dept("comp").artist = "Rahul"
        model = ShotTableModel([shot], user_role="artist", user_identities=["rahul"])
        mine = model.index(0, _col(model, "comp")).data(Qt.ItemDataRole.BackgroundRole)
        other = model.index(0, _col(model, "roto")).data(Qt.ItemDataRole.BackgroundRole)
        assert mine is not None and other is None


# ------------------------------------------------------------------ grouping
class TestGrouping:

    def _stack(self, shots):
        from slate.gui.tabs.vfx_dashboard_pro.ui.shot_proxy_models import ShotFilterProxy, ShotGroupModel
        model = ShotTableModel(shots, user_role="supervisor")
        proxy = ShotFilterProxy()
        proxy.setSourceModel(model)
        group = ShotGroupModel()
        group.setSourceModel(proxy)
        return model, proxy, group

    def test_groups_come_in_their_own_order_whatever_the_sort(self):
        """DSH-066/149/065."""
        shots = [_shot("SH1", "R05"), _shot("SH2", "R01"), _shot("SH3", "R02", assigned_artist="vikram singh")]
        model, proxy, group = self._stack(shots)
        group.set_group_by("Reel / Sequence")
        titles = [group.get_group_info(r)["title"] for r in group.get_header_rows()]
        assert titles == ["R01", "R02", "R05"]
        group.sort(_col(model, "shot_name"), Qt.SortOrder.DescendingOrder)
        titles = [group.get_group_info(r)["title"] for r in group.get_header_rows()]
        assert titles == ["R01", "R02", "R05"]
        group.set_group_by("Artist")
        assert "vikram singh" in [group.get_group_info(r)["title"] for r in group.get_header_rows()]
        group.set_group_by("Status")
        assert all(not group.get_group_info(r)["show_progress"] for r in group.get_header_rows())

    def test_group_header_colours_are_valid_and_hover_differs(self):
        """DSH-064/067."""
        from slate.gui.tabs.vfx_dashboard_pro.ui.group_header_delegate import GroupHeaderDelegate
        normal, hover = GroupHeaderDelegate.colours(False), GroupHeaderDelegate.colours(True)
        assert all(c.isValid() for c in normal.values())
        assert normal["background"] != hover["background"]
        assert normal["separator"] != normal["background"]


# ------------------------------------------------------------------ saving
class TestSaving:

    def test_a_save_writes_only_the_edited_shot(self, qtbot, mock_db):
        """DSH-018: one edit used to rewrite and version-bump every shot."""
        _project(mock_db, [_shot("SH010"), _shot("SH020")])
        widget = _open(_widget(qtbot))
        before = {s.shot_name: s.version for s in widget.all_shots}
        target = next(s for s in widget.all_shots if s.shot_name == "SH010")
        widget.table_model.apply_edit([target], lambda s: setattr(s, "sow", "new"), "sow")
        assert widget.save_btn.text() == "Save 1 change"
        assert widget.save_changes() is True
        after = {s.shot_name: s.version for s in widget.data_handler.read_shots()}
        assert after["SH020"] == before["SH020"]
        assert after["SH010"] == before["SH010"] + 1
        assert widget.has_unsaved_changes() is False

    def test_batch_saves_write_history_and_notify(self, qtbot, mock_db):
        """DSH-005/051."""
        _project(mock_db, [_shot("SH010"), _shot("SH020"), _shot("SH030")])
        widget = _open(_widget(qtbot))
        shots = widget.all_shots[:2]
        widget.on_batch_update(shots, {"status": "APPROVED", "assigned_artist": "Rahul"})
        assert widget.save_changes() is True
        history = mock_db.get_history(PROJECT, "SH010") or []
        fields = {row.get("field_changed") or row.get("field") for row in history}
        assert "status" in fields and "assigned_artist" in fields

    def test_a_shot_status_change_tells_the_shots_artist(self, mock_db):
        """A6: a verdict on a version with no department moved the shot's status and told nobody."""
        from types import SimpleNamespace
        handler = _project(mock_db, [_shot("SH010", assigned_artist="Rahul")])
        sent = []
        handler.notifier = SimpleNamespace(add_notification=lambda who, msg, kind: sent.append((who, msg)))
        shot = handler.read_shots()[0]
        shot.status = "RETAKE"
        assert handler.write_shots([shot])
        assert ("Rahul", "SH010 is now RETAKE.") in sent

    def test_a_conflict_names_the_shot_and_its_reel(self, mock_db):
        """DSH-019/112: structured conflicts."""
        handler = _project(mock_db, [_shot("SH010", "R01"), _shot("SH010", "R02")])
        mine = handler.read_shots()
        other = SQLiteHandler(PROJECT, db_manager=mock_db, user_role="supervisor", username="sup")
        theirs = other.read_shots()
        r02 = next(s for s in theirs if s.reel_episode == "R02")
        r02.sow = "theirs"
        assert other.write_shots([r02]) is True
        mine_r02 = next(s for s in mine if s.reel_episode == "R02")
        mine_r02.sow = "mine"
        with pytest.raises(StaleDataError) as caught:
            handler.write_shots([mine_r02])
        assert caught.value.conflicts[0]["reel"] == "R02"

    def test_a_refused_department_write_saves_nothing(self, mock_db, monkeypatch):
        """DSH-004 / DSH2-065: the shot row and its department rows are one transaction."""
        handler = _project(mock_db, [_shot("SH010")])
        shot = handler.read_shots()[0]
        shot.sow = "new scope"
        shot.dept("roto").status = "WIP"
        from slate.core.infra.db_results import DatabaseWriteError

        def refuse(*_a, **_k):
            raise DatabaseWriteError("refused")
        monkeypatch.setattr(handler, "_write_tasks", refuse)
        assert handler.write_shots([shot]) is False
        assert handler.last_error
        after = handler.read_shots()[0]
        assert after.sow == "" and after.dept("roto").status == ""

    def test_status_no_longer_mirrors_into_comp(self, mock_db):
        """DSH-023."""
        handler = _project(mock_db, [_shot("SH010")])
        shot = handler.read_shots()[0]
        shot.status = "APPROVED"
        assert handler.write_shots([shot])
        after = handler.read_shots()[0]
        assert after.status == "APPROVED" and after.dept("comp").status == ""

    def test_a_lead_cannot_change_shot_fields(self, mock_db):
        """DSH-020/021."""
        _project(mock_db, [_shot("SH010")])
        from tests.dashboard_util import person
        lead = SQLiteHandler(PROJECT, db_manager=mock_db, user_role=["lead"],
                             username=person(mock_db, "rl", "Roto Lead"))
        shot = lead.read_shots()[0]
        shot.status = "APPROVED"
        with pytest.raises(PermissionError):
            lead.write_shots([shot])
        nobody = SQLiteHandler(PROJECT, db_manager=mock_db, user_role=["lead"],
                               username=person(mock_db, "nl", "Trainee"))
        shot = nobody.read_shots()[0]
        shot.dept("roto").status = "WIP"
        with pytest.raises(PermissionError):
            nobody.write_shots([shot])


# ------------------------------------------------------------------ the screen
class TestScreen:

    def test_board_columns_are_real_statuses_and_a_drop_is_pending(self, qtbot, mock_db):
        """DSH-022/047/008."""
        _project(mock_db, [_shot("SH010", status="SENT FOR REVIEW"), _shot("SH020", status="READY")])
        widget = _open(_widget(qtbot))
        widget.board_btn.setChecked(True)
        assert "SENT FOR REVIEW" in widget.kanban_board.columns
        assert widget.kanban_board.columns["SENT FOR REVIEW"].count() == 1
        assert widget.kanban_board.columns["READY"].count() == 1
        task = next(i for i, s in enumerate(widget._board_shots) if s.shot_name == "SH010")
        widget.on_kanban_status_changed(task, "APPROVED")
        shot = widget._board_shots[task] if task < len(widget._board_shots) else None
        assert any(s.shot_name == "SH010" and s.status == "APPROVED" for s in widget.all_shots)
        assert widget.has_unsaved_changes()
        widget.undo_last_edit()
        assert not widget.has_unsaved_changes()

    def test_search_reads_every_column(self, qtbot, mock_db):
        """DSH-035."""
        _project(mock_db, [_shot("SH010", sow="remove rig", shot_type="2D Comp"), _shot("SH020")])
        widget = _open(_widget(qtbot))
        for text in ("remove rig", "  2d comp", "sh010"):
            widget.search_input.setText(text)
            widget.apply_filters()
            assert [s.shot_name for s in widget.displayed_shots] == ["SH010"], text

    def test_filters_do_not_follow_you_to_the_next_project(self, qtbot, mock_db):
        """DSH-036."""
        _project(mock_db, [_shot("SH010")])
        widget = _open(_widget(qtbot))
        widget.search_input.setText("zzz")
        widget.apply_filters()
        assert widget.displayed_shots == []
        assert widget.filter_chip.isVisibleTo(widget)
        widget.switch_project(PROJECT)
        assert len(widget.displayed_shots) == 1
        assert widget.active_filter_count() == 0

    def test_column_filter_values_come_from_every_shot(self, qtbot, mock_db):
        """DSH-037."""
        _project(mock_db, [_shot("SH1", edit_frames=1200), _shot("SH2", edit_frames=240),
                           _shot("SH3", shot_type="AI Shot")])
        widget = _open(_widget(qtbot))
        widget.header_view.apply_filter("type", {"AI Shot"})
        assert [s.shot_name for s in widget.displayed_shots] == ["SH3"]
        assert "(blank)" in widget.column_filter_values("type")
        assert widget.column_filter_values("frames")[:2] == ["240", "1200"]

    def test_stat_counters_filter_and_name_blank_status(self, qtbot, mock_db):
        """DSH-058/059/060/062."""
        _project(mock_db, [_shot("SH1", status="RETAKE"), _shot("SH2", status=""),
                           _shot("SH3", status="SENT FOR REVIEW")])
        widget = _widget(qtbot)
        assert not widget.stats_widget.isVisibleTo(widget)
        _open(widget)
        texts = [b.text() for b in widget.stats_widget.stat_containers.values()]
        assert "1 No status" in texts and "1 SENT FOR REVIEW" in texts
        widget.stats_widget.stat_containers["RETAKE"].click()
        assert [s.shot_name for s in widget.displayed_shots] == ["SH1"]
        widget.stats_widget.stat_containers["RETAKE"].click()
        assert len(widget.displayed_shots) == 3

    def test_an_artist_saves_their_own_status_and_can_undo(self, qtbot, mock_db):
        """DSH-047: an artist's own status saves at once, with Undo."""
        shot = _shot("SH010")
        shot.dept("comp").artist = "Rahul"
        _project(mock_db, [shot])
        widget = _open(_widget(qtbot, role="Artist", username="rahul", display_name="Rahul"))
        assert not widget.save_btn.isVisibleTo(widget)
        model = widget.table_model
        assert model.setData(model.index(0, _col(model, "comp")), "WIP") is True
        stored = widget.data_handler.read_shots()[0]
        assert stored.dept("comp").status == "WIP"
        widget.undo_last_edit()
        assert widget.data_handler.read_shots()[0].dept("comp").status == ""

    def test_scopes_have_proper_names_for_everyone(self, qtbot, mock_db):
        """DSH-033/126/127/128."""
        _project(mock_db, [_shot()])
        widget = _open(_widget(qtbot, role="Developer"))
        items = [widget.scope_combo.itemText(i) for i in range(widget.scope_combo.count())]
        assert "Matte Painting department" in items
        assert "My shots" not in items           # a developer is not given work
        artist = _open(_widget(qtbot, role="Artist", username="rahul", display_name="Rahul"))
        items = [artist.scope_combo.itemText(i) for i in range(artist.scope_combo.count())]
        assert "All shots" not in items

    def test_reel_and_shot_name_cannot_be_hidden(self, qtbot, mock_db):
        """DSH-132."""
        _project(mock_db, [_shot()])
        widget = _open(_widget(qtbot))
        widget._set_all_columns(False)
        assert not widget.table.isColumnHidden(_col(widget.table_model, "shot_name"))
        assert not widget.table.isColumnHidden(_col(widget.table_model, "reel"))

    def test_the_last_project_is_remembered(self, qtbot, mock_db):
        """DSH-105."""
        _project(mock_db, [_shot()])
        _open(_widget(qtbot))
        again = _widget(qtbot)
        assert again.current_project is not None and again.current_project.code == PROJECT

    def test_detail_panel_save_changes_only_what_was_touched(self, qtbot, mock_db):
        """DSH-002/090/034: unknown values survive; Apply stages only the change."""
        shot = _shot("SH010", status="CBB", shot_type="Roto only", target="TBD")
        shot.dept("roto").status = "SENT FOR REVIEW"
        shot.dept("prep").status = "Done"
        _project(mock_db, [shot])
        widget = _open(_widget(qtbot))
        loaded = widget.all_shots[0]
        widget.open_detail_dock(loaded)
        panel = widget.detail_widget
        assert panel.collect_changes() == {}
        panel.apply_changes()
        assert widget.has_unsaved_changes() is False
        panel.sow_edit.setPlainText("new scope")
        panel.apply_changes()
        assert loaded.sow == "new scope"
        assert loaded.status == "CBB" and loaded.shot_type == "Roto only" and loaded.target == "TBD"
        assert loaded.dept("roto").status == "SENT FOR REVIEW" and loaded.dept("prep").status == "Done"
        assert widget.has_unsaved_changes() is True

    def test_detail_panel_is_read_only_for_a_compositor(self, qtbot, mock_db):
        """DSH-027."""
        shot = _shot("SH010")
        shot.dept("comp").artist = "Priya Sharma"
        _project(mock_db, [shot])
        widget = _open(_widget(qtbot, role="Compositor", username="priya", display_name="Priya Sharma"))
        widget.open_detail_dock(widget.all_shots[0])
        panel = widget.detail_widget
        assert not panel.status_combo.isEnabled()
        assert panel.sow_edit.isReadOnly()
        assert panel.depts["comp"]["status_combo"].isEnabled()
        assert not panel.depts["roto"]["status_combo"].isEnabled()

    def test_linked_heroes_can_be_removed(self, qtbot, mock_db):
        """DSH-091."""
        _project(mock_db, [_shot("SH010", similar_to=["HERO1"]), _shot("HERO1", is_hero=True)])
        widget = _open(_widget(qtbot))
        shot = next(s for s in widget.all_shots if s.shot_name == "SH010")
        widget.open_detail_dock(shot)
        panel = widget.detail_widget
        panel.similar_list.setCurrentRow(0)
        panel._remove_similar()
        panel.apply_changes()
        assert shot.similar_to == []


# ------------------------------------------------------------------ projects
class TestProjects:

    def test_archive_hides_and_restores_without_losing_history(self, mock_db):
        """DSH-053."""
        from slate.gui.tabs.vfx_dashboard_pro.core.project_manager import ProjectManager
        _project(mock_db, [_shot()])
        manager = ProjectManager()
        assert manager.archive_project(PROJECT, roles=["coordinator"]) is False
        assert manager.archive_project(PROJECT, roles=["admin"], by="admin") is True
        assert PROJECT not in ProjectManager().projects
        assert mock_db.get_tracking_shots(PROJECT)
        assert {"code": PROJECT, "name": "Dashboard test"} in manager.archived_projects()
        assert manager.restore_project(PROJECT, roles=["admin"]) is True
        assert PROJECT in ProjectManager().projects

    def test_delete_keeps_the_history(self, mock_db):
        from slate.gui.tabs.vfx_dashboard_pro.core.project_manager import ProjectManager
        _project(mock_db, [_shot()])
        manager = ProjectManager()
        assert manager.delete_project(PROJECT, roles=["admin"], by="admin") is True
        assert not mock_db.get_tracking_shots(PROJECT)
        assert mock_db.get_history(PROJECT)

    def test_project_root_is_saved_to_the_database(self, mock_db, tmp_path):
        """DSH-031."""
        from slate.gui.tabs.vfx_dashboard_pro.core.project_manager import ProjectManager
        _project(mock_db, [_shot()])
        manager = ProjectManager()
        assert manager.set_project_folder_base(PROJECT, str(tmp_path)) is True
        config = ProjectManager().get_project(PROJECT)
        assert config.folder_base == str(tmp_path)
        assert config.folder_template["roto"].endswith("04_Roto")
        assert "output" in config.folder_template


# ------------------------------------------------------------------ more
class TestMore:

    def test_actual_days_are_edited_in_the_panel_and_stored(self, qtbot, mock_db):
        """Production: tracking_tasks.actual_days, editable per department."""
        mock_db.execute_update("ALTER TABLE tracking_tasks ADD COLUMN actual_days REAL")
        shot = _shot("SH010")
        shot.dept("comp").bid_days = 3.0
        _project(mock_db, [shot])
        widget = _open(_widget(qtbot))
        loaded = widget.all_shots[0]
        widget.open_detail_dock(loaded)
        widget.detail_widget.depts["comp"]["actual_spin"].setValue(4.5)
        widget.detail_widget.apply_changes()
        assert widget.save_changes() is True
        row = mock_db.execute_query(
            "SELECT actual_days FROM tracking_tasks WHERE shot_id=%s AND department='comp'",
            (loaded.id,), fetch="one")
        assert float(row["actual_days"]) == 4.5
        assert widget.data_handler.read_shots()[0].dept("comp").actual_days == 4.5

    def test_the_passbook_keeps_one_row_per_reel_and_numbers_it(self, tmp_path):
        """DSH-101."""
        from openpyxl import load_workbook
        from slate.gui.tabs.vfx_dashboard_pro.core.excel_handler import ExcelHandler
        from slate.gui.tabs.vfx_dashboard_pro.core.project_manager import ProjectConfig, default_column_mapping
        path = tmp_path / "book.xlsx"
        config = ProjectConfig(code="P", name="P", excel_path=str(path), column_mapping=default_column_mapping())
        handler = ExcelHandler(str(path), config)
        assert handler.create_workbook()
        a, b = _shot("SH010", "R01"), _shot("SH010", "R02")
        a.thumbnail_path = r"C:\Users\x\AppData\Local\Slate\cache\placeholder_red.png"
        assert ExcelHandler(str(path), config).write_shots([a, b])
        ws = load_workbook(path)[config.sheet_name] if config.sheet_name in load_workbook(path).sheetnames \
            else load_workbook(path).active
        rows = [r for r in ws.iter_rows(min_row=3, values_only=True) if r[3]]
        assert len(rows) == 2
        assert [r[0] for r in rows] == [1, 2]
        assert all(not r[1] for r in rows)

    def test_smart_search_is_gone(self):
        """DSH-082 / FIX_PLAN: no hidden '?' search, no model download."""
        import importlib.util
        assert importlib.util.find_spec("slate.core.domain.vector_service") is None
        assert not hasattr(Shot(), "_semantic_embedding")

    def test_batch_edit_uses_the_studio_lists(self, qtbot):
        """DSH-111/120."""
        from slate.core.domain import shot_status
        from slate.gui.tabs.vfx_dashboard_pro.ui.batch_edit_dialog import BatchEditDialog
        dialog = BatchEditDialog(3, all_users=["Rahul"])
        qtbot.addWidget(dialog)
        types = [dialog.type_combo.itemText(i) for i in range(dialog.type_combo.count())]
        assert types == shot_status.shot_types()
        dialog.priority_cb.setChecked(True)
        dialog.priority_combo.setCurrentIndex(dialog.priority_combo.findData(0))
        assert dialog.get_updates() == {"priority": 0}


def test_placeholder_thumbnail_paths_are_cleared_once(mock_db):
    """DSH-052: placeholder paths already stored are repaired; real ones kept."""
    from slate.core.infra.migrations.dashboard_repairs import clear_placeholder_thumbnails
    a = _shot("SH010", thumbnail_path=r"C:\Users\dev\Slate\V0040\slate\gui\tabs\vfx_dashboard_pro\cache\thumbnails\placeholder_red.png")
    b = _shot("SH020", thumbnail_path=r"\server\thumbs\SH020.jpg")
    handler = _project(mock_db, [a, b])
    assert clear_placeholder_thumbnails(mock_db) is True
    paths = {s.shot_name: s.thumbnail_path for s in handler.read_shots()}
    assert paths == {"SH010": "", "SH020": r"\server\thumbs\SH020.jpg"}


def test_column_widths_survive_showing_the_grid(qtbot, mock_db):
    """The frozen overlay's hidden columns must not shrink the grid's (seen on screen)."""
    _project(mock_db, [_shot("SH010")])
    widget = _open(_widget(qtbot))
    before = [widget.table.columnWidth(i) for i in range(6)]
    widget.resize(1200, 700)
    widget.show()
    qtbot.wait(200)
    assert [widget.table.columnWidth(i) for i in range(6)] == before


def test_row_colours_are_real_colours():
    """A modified row was painted black: QColor('rgba(...)') is invalid in Qt."""
    shot = _shot()
    model = ShotTableModel([shot], user_role="supervisor")
    model.setData(model.index(0, _col(model, "sow")), "x")
    colour = model.index(0, 0).data(Qt.ItemDataRole.BackgroundRole)
    assert colour.isValid() and colour.alpha() < 255


# ------------------------------------------------------------------ verifier round
class TestVerifierRound:

    def test_delivery_batches_open_from_reports(self, qtbot, mock_db, monkeypatch):
        """DSH-029/118: the dialog opened from the dashboard used to raise TypeError."""
        from slate.gui.tabs.vfx_dashboard_pro.ui import delivery_batches_dialog as dbd
        _project(mock_db, [_shot()])
        artist = _open(_widget(qtbot, role="Artist", username="rahul", display_name="Rahul"))
        opened = []
        monkeypatch.setattr(dbd.DeliveryBatchesDialog, "exec", lambda self: opened.append(self) or 0)
        artist.open_delivery_batches_dialog()
        # DSH2-035: an artist, who sees only their own shots, does not get the
        # project's packages at all; the store still refuses them a create.
        assert opened == []
        dialog = dbd.DeliveryBatchesDialog(PROJECT, roles=artist.access_roles)
        qtbot.addWidget(dialog)
        assert not dialog.create_btn.isVisibleTo(dialog)
        with pytest.raises(PermissionError):
            dialog.store.create_delivery(PROJECT, "DEL", [1])
        sup = _open(_widget(qtbot))
        sup.open_delivery_batches_dialog()
        assert opened[0].store.roles == sup.access_roles
        assert opened[0].create_btn.isVisibleTo(opened[0])

    def test_long_project_names_end_in_an_ellipsis(self, qtbot):
        """DSH-055."""
        from slate.gui.tabs.vfx_dashboard_pro.ui.eliding_combo import ElidingComboBox
        combo = ElidingComboBox()
        qtbot.addWidget(combo)
        combo.addItem("KLC - Kaalchakra – The Wheel of Time (कालचक्र) Season 1 Final Delivery")
        combo.setFixedWidth(200)
        combo.show()
        assert combo.elided_text().endswith("…")
        assert combo.toolTip() == combo.currentText() or combo.toolTip() == ""

    def test_counters_fit_and_do_not_double_count(self, qtbot):
        """DSH-061: the row collapsed to '262 shots +14 more' at 1280."""
        from slate.gui.tabs.vfx_dashboard_pro.ui.stats_widget import StatsWidget
        stats = StatsWidget()
        qtbot.addWidget(stats)
        shots = [_shot(f"S{i}", status=s) for i, s in enumerate(
            ["YTS", "READY", "WIP", "SENT FOR REVIEW", "RETAKE", "APPROVED", "OMIT", ""] * 30)]
        stats.resize(900, 30)
        stats.show()
        stats.update_stats(shots)
        qtbot.wait(50)
        hidden = stats._hidden
        assert len(hidden) == len(set(hidden))
        assert len(stats.visible_texts()) >= 5
        stats.resize(2000, 30)
        qtbot.wait(50)
        assert stats._hidden == []

    def test_publish_messages_name_the_real_folder(self, qtbot, mock_db, tmp_path):
        """DSH-095: new projects deliver to 08_Deliver, not '08_Output'."""
        _project(mock_db, [_shot()])
        widget = _open(_widget(qtbot))
        widget.project_manager.set_project_folder_base(PROJECT, str(tmp_path))
        widget.current_project = widget.project_manager.get_project(PROJECT)
        assert widget.output_folder_name(widget.all_shots[0]) == "08_Deliver"

    def test_board_card_text_is_whole_lines(self):
        """NEW-dashboard-1: SOW on a card is two whole lines, the second elided."""
        from PySide6.QtGui import QFont, QFontMetrics
        from slate.gui.widgets.kanban_board import KanbanCard
        fm = QFontMetrics(QFont())
        text = KanbanCard.two_lines("remove rig and add sky " * 20, fm, 150)
        lines = text.split("\n")
        assert len(lines) == 2 and lines[1].endswith("…")
        assert all(fm.horizontalAdvance(l) <= 150 for l in lines)

    def test_frozen_overlay_has_no_frame_across_headings(self, qtbot, mock_db):
        """NEW-dashboard-2."""
        from PySide6.QtWidgets import QFrame
        _project(mock_db, [_shot("A", "R01"), _shot("B", "R02")])
        widget = _open(_widget(qtbot))
        widget.groupby_combo.setCurrentIndex(widget.groupby_combo.findData("Reel / Sequence"))
        assert widget.table.frozen.frameShape() == QFrame.Shape.NoFrame
        widget.resize(1200, 600)
        widget.show()
        qtbot.wait(100)
        widget.table.frozen.viewport().repaint()

    def test_a_producer_gets_the_dashboard_tab(self, mock_db):
        """NEW-dashboard-3: new installs and upgrades agree."""
        import json
        from slate.core.domain.permissions_catalog import STUDIO_ROLES
        from slate.core.domain.user_manager import UserManager
        assert "Dashboard" in STUDIO_ROLES["Producer"]
        mock_db.execute_update("DELETE FROM ut_roles WHERE LOWER(role_name)='producer'")
        mock_db.execute_update("INSERT INTO ut_roles (role_name, permissions) VALUES (%s, %s)",
                               ("Producer", json.dumps(["Reports", "Settings"])))
        mock_db.execute_update("CREATE TABLE IF NOT EXISTS ut_role_upgrades (name TEXT PRIMARY KEY, applied_at TEXT)")
        mock_db.execute_update("DELETE FROM ut_role_upgrades WHERE name='2026-10-producer-dashboard'")
        manager = UserManager.__new__(UserManager)
        manager._get_db = lambda: mock_db
        manager._upgrade_producer_dashboard()
        row = mock_db.execute_query("SELECT permissions FROM ut_roles WHERE role_name='Producer'", fetch="one")
        perms = json.loads(row["permissions"])
        assert "Dashboard" in perms and "Reports" in perms

    def test_an_existing_studio_gets_the_new_role_defaults_once(self, mock_db):
        """IT, Production Coordinator and Producer get what a new database gives them; only once."""
        import json
        from slate.core.domain.user_manager import UserManager
        old = {"IT": ["IT", "Settings"], "Production Coordinator": ["Dashboard"],
               "Producer": ["Dashboard"], "Artist": ["Dashboard"]}
        for role, perms in old.items():
            mock_db.execute_update("DELETE FROM ut_roles WHERE role_name=%s", (role,))
            mock_db.execute_update("INSERT INTO ut_roles (role_name, permissions) VALUES (%s, %s)",
                                   (role, json.dumps(perms)))
        mock_db.execute_update("CREATE TABLE IF NOT EXISTS ut_role_upgrades (name TEXT PRIMARY KEY, applied_at TEXT)")
        mock_db.execute_update("DELETE FROM ut_role_upgrades WHERE name='2026-10-new-defaults'")
        manager = UserManager.__new__(UserManager)
        manager._get_db = lambda: mock_db

        def perms(role):
            row = mock_db.execute_query("SELECT permissions FROM ut_roles WHERE role_name=%s", (role,), fetch="one")
            return json.loads(row["permissions"])

        manager._upgrade_new_defaults()
        assert {"Admin Panel", "can:manage_users", "can:fleet_control"} <= set(perms("IT"))
        assert {"Bidding", "can:bid_write"} <= set(perms("Producer")) & set(perms("Production Coordinator"))
        assert perms("Artist") == ["Dashboard"], "other roles untouched"
        # The studio takes Bidding away again: the upgrade does not come back.
        mock_db.execute_update("UPDATE ut_roles SET permissions=%s WHERE role_name='Producer'",
                               (json.dumps(["Dashboard"]),))
        manager._upgrade_new_defaults()
        assert perms("Producer") == ["Dashboard"]


class TestShowMyShotsFromHome:
    """Integration: Home's 'See all my shots' uses the dashboard's own 'my shots' scope."""

    @pytest.mark.parametrize("role", ["Compositor", "Supervisor", "Artist"])
    def test_show_my_shots_selects_the_scope_and_clears_search(self, qtbot, mock_db, role):
        _project(mock_db, [_shot("SH010", assigned_artist="priya"),
                           _shot("SH020", assigned_artist="Priya Sharma"),
                           _shot("SH030", assigned_artist="rahul")])
        widget = _open(_widget(qtbot, role=role, username="priya", display_name="Priya Sharma"))
        widget.search_input.setText("SH0")
        from slate.core.domain.access import can_be_assigned
        if not widget._is_artist_scope() and not can_be_assigned([role.lower()]):
            pytest.skip("this role is never given shots")
        assert widget.show_my_shots()
        assert widget.search_input.text() == ""
        assert widget.scope_combo.currentData() in ("my_shots", "all")
        assert sorted(s.shot_name for s in widget.displayed_shots) == ["SH010", "SH020"]

    def test_home_calls_the_scope(self, qtbot):
        from slate.gui.tabs.home_tab import HomeTab
        calls = []

        class Dash:
            search_input = None

            def show_my_shots(self):
                calls.append("scope")
                return True

        class Host:
            def _switch_to_tab_label(self, label):
                return label == "VFX Dashboard"

            def _get_tab_instance(self, label, create=False):
                return Dash()

        home = HomeTab(user_data={"username": "priya", "display_name": "Priya Sharma"}, mode="vfx")
        qtbot.addWidget(home)
        home._host = lambda: Host()
        assert home.see_all_shots() and calls == ["scope"]
