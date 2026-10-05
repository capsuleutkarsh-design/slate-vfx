"""
The Attendance tab, offscreen, on a database of its own.

    the punch buttons offer only the valid next step                  HR-021
    the Edit punch dialog checks before it closes, refuses out < in    HR-017/018
    a supervisor sees their reports only, read-only                   HR-013
    a very long name does not take the grid's width                   HR-014
    search, filters and name order on the grid                        HR-015
    the personal view can show any month                              HR-022
    one colour map for the legend and the tables                      HR-028
    the export agrees with the grid (overnight, codes, '01 Tue')      HR-010/042
"""

from datetime import date, datetime, timedelta

import pytest

pytest.importorskip("PySide6.QtWidgets")

from PySide6.QtCore import QTime  # noqa: E402
from PySide6.QtWidgets import QApplication  # noqa: E402

from slate.core.domain import leave_policy as lp  # noqa: E402
from slate.core.domain.central_attendance import CentralAttendance  # noqa: E402


@pytest.fixture(scope="module")
def app():
    return QApplication.instance() or QApplication([])


@pytest.fixture(autouse=True)
def _tidy(app):
    yield
    import gc
    for widget in QApplication.topLevelWidgets():
        widget.close()
        widget.deleteLater()
    QApplication.processEvents()
    gc.collect()
    QApplication.processEvents()


class _Context:
    def attendance(self):
        raise AssertionError("passed explicitly")

    def user_manager(self):
        raise AssertionError("passed explicitly")


LONG = "Venkatanarasimharajuvaripeta Subrahmanyam Krishnamurthy Iyengar"


@pytest.fixture
def studio(mock_db, app):
    from slate.core.domain import access, people
    from slate.core.domain.user_manager import UserManager
    access.reset_cache()
    people.refresh()
    lp.set_overrides({})
    um = UserManager(db=mock_db)
    um.add_user("hr.meera", "pw", ["HR"], "Meera Iyer", "HR", location="Mumbai")
    um.add_user("sup.vikram", "pw", ["Supervisor"], "Vikram Malhotra", "Comp", location="Mumbai")
    um.add_user("aarav", "pw", ["Artist"], "Aarav Sharma", "Comp", reports_to="sup.vikram",
                location="Mumbai")
    um.add_user("diya", "pw", ["Artist"], "Diya Nair", "Roto", reports_to="hr.meera",
                location="Chennai")
    um.add_user("venkat", "pw", ["Artist"], LONG, "Comp", reports_to="sup.vikram",
                location="Chennai")
    people.refresh()
    yield {"db": mock_db, "um": um, "att": CentralAttendance(mock_db)}
    lp.set_overrides({})
    access.reset_cache()
    people.refresh()


def _tab(studio, username, roles, display=""):
    from slate.gui.attendance_tab import AttendanceTab
    user = {"user_id": username, "username": username, "display_name": display or username,
            "roles": roles}
    return AttendanceTab(user, attendance=studio["att"], user_manager=studio["um"],
                         app_context=_Context())


def _names(tab):
    t = tab.team_table
    # The full name is in the tooltip ("Name (username)"); the header text is cut.
    return [t.verticalHeaderItem(r).toolTip().rsplit(" (", 1)[0] for r in range(t.rowCount())]


# ----------------------------------------------------------------- HR-021

def _punch(tab, qtbot, action):
    """The punch runs on a worker (CentralAttendance.punch); wait for it."""
    tab.manual_punch(action)
    qtbot.waitUntil(lambda: not tab.punch_busy(), timeout=5000)


def test_the_buttons_follow_the_day(studio, qtbot):
    tab = _tab(studio, "aarav", ["Artist"])
    assert tab.btn_punch_in.isEnabled() and not tab.btn_punch_out.isEnabled()
    _punch(tab, qtbot, "in")
    assert not tab.btn_punch_in.isEnabled() and tab.btn_punch_out.isEnabled()
    assert tab.lbl_punch_note.text().startswith("Punched in at")
    assert tab.lbl_status.text().startswith("Working")
    _punch(tab, qtbot, "out")                     # the question is answered Yes in tests
    assert tab.btn_punch_in.isEnabled() and tab.btn_punch_in.text() == "Punch in again"
    assert not tab.btn_punch_out.isEnabled()


# ------------------------------------------------------------- HR-017 / 018

def test_the_edit_dialog_refuses_out_before_in_and_stays_open():
    from slate.gui.attendance_tab import EditPunchDialog
    dialog = EditPunchDialog("Aarav", date(2026, 9, 15), "18:00", "09:00")
    dialog.reason.setText("fix")
    assert "before the in time" in dialog.problem()
    dialog._save()
    assert not dialog.error.isHidden() and dialog.result() == 0
    dialog.overnight.setChecked(True)
    assert dialog.problem() == ""


def test_the_edit_dialog_needs_a_time_and_a_reason():
    from slate.gui.attendance_tab import EditPunchDialog
    dialog = EditPunchDialog("Aarav", date(2026, 9, 15))
    # A day with no record opens with both times ready to type (HR2-016).
    assert not dialog.no_in.isChecked() and not dialog.no_out.isChecked()
    assert dialog.e_in.isEnabled() and dialog.e_out.isEnabled()
    dialog.no_in.setChecked(True)
    dialog.no_out.setChecked(True)
    assert "Enter an in time" in dialog.problem()
    dialog.no_in.setChecked(False)
    dialog.e_in.setTime(QTime(9, 30))
    assert "Say why" in dialog.problem()


# ----------------------------------------------------------------- HR-013

def test_a_supervisor_sees_only_their_reports_read_only(studio):
    tab = _tab(studio, "sup.vikram", ["Supervisor"])
    assert tab.team_table is not None
    assert sorted(_names(tab)) == sorted(["Aarav Sharma", LONG])
    assert tab.btn_import is None and tab.btn_holidays is None
    called = []
    tab.show_edit_dialog = lambda *a, **k: called.append(a)
    tab.on_cell_double_click(0, len(tab.TEAM_STATS))
    assert called == [], "supervisors do not correct punches"


def test_hr_sees_everybody_sorted_by_name_without_system_accounts(studio):
    tab = _tab(studio, "hr.meera", ["HR"])
    names = _names(tab)
    assert names == sorted(names, key=str.casefold)
    assert "System Admin" not in names and "QA Tester" not in names
    tab.show_system.setChecked(True)
    assert "System Admin" in _names(tab)


# ----------------------------------------------------------------- HR-014

def test_a_long_name_does_not_take_the_grid(studio):
    tab = _tab(studio, "hr.meera", ["HR"])
    assert tab.team_table.verticalHeader().width() == 200
    row = _names(tab).index(LONG)
    assert LONG in tab.team_table.verticalHeaderItem(row).toolTip()


# ----------------------------------------------------------------- HR-015

def test_search_and_location_filter(studio):
    tab = _tab(studio, "hr.meera", ["HR"])
    reads = []
    real = tab.attendance.get_full_month_data
    tab.attendance.get_full_month_data = lambda *a: reads.append(a) or real(*a)
    tab.team_search.setText("a")
    tab.team_search.setText("aarav")
    assert _names(tab) != ["Aarav Sharma"], "the search waits for the typing to stop (HR2-020)"
    assert tab._search_timer.isActive() and tab._search_timer.interval() == 300
    tab._search_timer.stop()
    tab._search_timer.timeout.emit()              # the pause after the last key
    assert _names(tab) == ["Aarav Sharma"]
    assert reads == [], "the people already read are filtered, not read again"
    tab.team_search.setText("")
    tab._search_timer.stop()
    tab._search_timer.timeout.emit()
    tab.filter_location.setCurrentIndex(tab.filter_location.findData("Chennai"))
    assert sorted(_names(tab)) == sorted(["Diya Nair", LONG])


# ----------------------------------------------------------------- HR-022

def test_the_personal_view_shows_any_month(studio):
    att = studio["att"]
    last_month = (date.today().replace(day=1) - timedelta(days=1))
    day = last_month.replace(day=10)
    att.write_day("aarav", day, "09:30", "18:30")
    tab = _tab(studio, "aarav", ["Artist"])
    tab.my_month.setCurrentIndex(last_month.month - 1)
    tab.my_year.setValue(last_month.year)
    assert tab.my_table.rowCount() == last_month.day
    assert tab.my_table.item(9, 1).text() == "09:30"


# ----------------------------------------------------------------- HR-028

def test_the_legend_is_built_from_the_colours_the_tables_use():
    from slate.gui.attendance_tab import legend_entries, state_styles
    used = {colour for colour, _ in state_styles().values()}
    for colour, label in legend_entries():
        assert colour in used and label


def test_streak_and_status_wording():
    from slate.gui.attendance_tab import status_text, streak_text
    assert streak_text(0) == ""
    assert streak_text(1) == "1 day on time"
    assert streak_text(5) == "5 days on time in a row"
    assert status_text({"state": "out"}) == "Not punched in today"
    assert status_text({"state": "done", "out": "18:30"}) == "Punched out at 18:30"


# ----------------------------------------------------------- HR-010 / 042

def test_the_export_agrees_with_the_grid(tmp_path):
    from openpyxl import load_workbook
    from slate.gui.attendance_export_worker import build_workbook, summarise
    lp.set_overrides({})
    data = {"aarav": {"01": {"in": "20:00", "out": "05:30"}, "13": {"in": "12:00", "out": "16:00"}}}
    rows = [{"username": "aarav", "name": "=HYPERLINK(\"x\")", "holidays": set()}]
    summary = summarise(rows[0], data["aarav"], 2026, 9, now=datetime(2026, 9, 30, 12))
    assert summary["days"][0][1] == pytest.approx(9.5)
    # Tuesday 1 Sep at 20:00 is late; Sunday 13 Sep at 12:00 is not.
    assert summary["late"] == 1

    path = tmp_path / "a.xlsx"
    build_workbook(str(path), 2026, 9, rows, data, now=datetime(2026, 9, 30, 12))
    ws = load_workbook(path).active
    headers = [c.value for c in ws[1]]
    assert "01 Tue" in headers
    assert str(ws.cell(row=2, column=1).value).startswith("'"), "formula neutralised"
    first_day = headers.index("01 Tue") + 1
    assert ws.cell(row=2, column=first_day).value == "L 9.5"      # late, 9.5 hours


# ------------------------------------------------------------- round 2

def test_a_supervisor_sees_the_people_under_their_leads(studio):
    """HR2-022: the reporting tree, not one level; HR2-031: no filters that filter nothing."""
    um = studio["um"]
    um.add_user("lead.kiran", "pw", ["Supervisor"], "Kiran Lead", "Comp",
                reports_to="sup.vikram", location="Mumbai")
    um.add_user("ira", "pw", ["Artist"], "Ira Rao", "Comp", reports_to="lead.kiran",
                location="Mumbai")
    from slate.core.domain import people
    people.refresh()
    tab = _tab(studio, "sup.vikram", ["Supervisor"])
    assert "Ira Rao" in _names(tab) and "Kiran Lead" in _names(tab)
    assert tab.show_system.isHidden()


def test_the_wfh_tick_waits_for_the_punch_in(studio, qtbot):
    """HR2-011: the refresh keeps a choice made before punching in."""
    tab = _tab(studio, "aarav", ["Artist"])
    tab.chk_wfh_box.setChecked(True)
    tab.refresh_personal_view()
    assert tab.chk_wfh_box.isChecked()
    _punch(tab, qtbot, "in")
    assert studio["att"].today_state("aarav")["wfh"] is True


def test_team_viewers_get_two_tabs_and_admin_no_punch_card(studio):
    """HR2-008 / HR2-029."""
    hr = _tab(studio, "hr.meera", ["HR"])
    assert hr.pages is not None and hr.pages.count() == 2
    artist = _tab(studio, "aarav", ["Artist"])
    assert artist.pages is None
    admin = _tab(studio, "admin", ["Admin"])
    assert admin.personal_card.isHidden() and admin.pages.count() == 1


def test_the_edit_dialog_shows_every_session():
    """HR2-001 / HR2-015 / HR2-016."""
    from slate.gui.attendance_tab import EditPunchDialog
    dialog = EditPunchDialog("Vihaan", date(2026, 9, 8), "09:00", "19:00", has_record=True,
                             sessions=[("09:00", "13:00"), ("17:00", "19:00")])
    dialog.reason.setText("check")
    assert dialog.values()["sessions"] == [("09:00", "13:00"), ("17:00", "19:00")]
    assert dialog.no_in.isHidden() and dialog.problem() == ""
    dialog._add_row()
    assert len(dialog.values()["sessions"]) == 3
    dialog._remove_row(dialog._rows[-1])
    dialog._remove_row(dialog._rows[-1])
    assert dialog.values()["sessions"] == [("09:00", "13:00")]
    auto = EditPunchDialog("Aarav", date(2026, 9, 15), "11:30", "19:30", has_record=True,
                           auto_cutoff="19:30:00", leave={"type": "Casual", "half": ""})
    texts = " ".join(w.text() for w in auto.findChildren(type(auto.error)))
    assert "automatic punch-out at 19:30" in texts and "Approved leave: Casual" in texts
