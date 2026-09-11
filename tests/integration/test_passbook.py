"""
The Excel passbook: the software writes it, in the central folder, from nothing.

A project made by Build & Ingest has no Excel path and no column mapping.
Before this, every save after an ingest reported "Excel mirror failed" and no
passbook was ever opened. Now the first save creates it where the software and
the server already exchange data, and every save after adds or updates rows.

Nothing here reads a sheet back into the database, and nothing ever will.
"""

import pytest
from openpyxl import load_workbook

from ut_vfx.core.domain.shot_registry import register_ingested_shots
from ut_vfx.core.infra.global_config import GlobalConfig
from ut_vfx.gui.tabs.vfx_dashboard_pro.core.project_manager import (
    ProjectManager, passbook_path,
)
from ut_vfx.gui.tabs.vfx_dashboard_pro.core.sqlite_handler import SQLiteHandler
from ut_vfx.gui.tabs.vfx_dashboard_pro.models.shot_model import Shot
from ut_vfx.gui.tabs.vfx_dashboard_pro.ui.dashboard_sync_service import (
    DashboardSyncService,
)


@pytest.fixture
def central(tmp_path, monkeypatch):
    """The central server folder, pointed at somewhere disposable."""
    folder = tmp_path / "central"
    folder.mkdir()
    monkeypatch.setattr(GlobalConfig, "server_root", classmethod(lambda cls: folder))
    return folder


def _ingested_project(mock_db, tmp_path, code="PRJ"):
    register_ingested_shots(
        code,
        [{"reel": "ReelA", "shot": "SH010", "scan_version": "v001"},
         {"reel": "ReelA", "shot": "SH020", "scan_version": "v001"}],
        db=mock_db, project_name="Passbook Test",
        folder_base=str(tmp_path / code),
    )
    return ProjectManager()


def _sheet(path):
    wb = load_workbook(path)
    return wb[wb.sheetnames[0]]


def _rows_by_shot(ws, header_row=2):
    """{shot name: {header: value}} for every data row."""
    headers = [c.value for c in ws[header_row]]
    rows = {}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        record = dict(zip(headers, row))
        name = record.get("SHOT NAME")
        if name:
            rows[name] = record
    return rows


class TestWhereItLives:

    def test_in_a_tracking_folder_under_the_central_root(self, central):
        assert passbook_path("PRJ") == central / "Tracking" / "PRJ_tracking.xlsx"

    def test_the_project_code_is_made_safe_for_a_filename(self, central):
        assert passbook_path("EP 01/Reel?") .name == "EP_01_Reel__tracking.xlsx"


class TestAnIngestedProjectGetsAPassbook:

    def test_the_first_save_creates_it(self, central, mock_db, tmp_path):
        manager = _ingested_project(mock_db, tmp_path)
        handler = SQLiteHandler("PRJ", db_manager=mock_db, user_role="supervisor")
        service = DashboardSyncService(manager)

        shots = handler.read_shots()
        ok, _ = service.mirror_shots_to_excel(
            shots, manager.get_project("PRJ"), handler, force=True)

        assert ok, service.last_backup_error
        assert (central / "Tracking" / "PRJ_tracking.xlsx").exists()

    def test_the_project_remembers_where_its_passbook_is(self, central, mock_db, tmp_path):
        manager = _ingested_project(mock_db, tmp_path)
        handler = SQLiteHandler("PRJ", db_manager=mock_db, user_role="supervisor")
        DashboardSyncService(manager).mirror_shots_to_excel(
            handler.read_shots(), manager.get_project("PRJ"), handler, force=True)

        reloaded = ProjectManager().get_project("PRJ")

        assert reloaded.excel_path == str(central / "Tracking" / "PRJ_tracking.xlsx")
        assert reloaded.column_mapping, "a sheet with no mapping has nowhere to put anything"

    def test_every_shot_gets_a_row(self, central, mock_db, tmp_path):
        """New shots used to be skipped: a passbook never gained a row."""
        manager = _ingested_project(mock_db, tmp_path)
        handler = SQLiteHandler("PRJ", db_manager=mock_db, user_role="supervisor")
        DashboardSyncService(manager).mirror_shots_to_excel(
            handler.read_shots(), manager.get_project("PRJ"), handler, force=True)

        rows = _rows_by_shot(_sheet(central / "Tracking" / "PRJ_tracking.xlsx"))

        assert set(rows) == {"SH010", "SH020"}
        assert rows["SH010"]["REEL"] == "ReelA"

    def test_a_later_save_updates_the_row_rather_than_adding_another(
            self, central, mock_db, tmp_path):
        manager = _ingested_project(mock_db, tmp_path)
        handler = SQLiteHandler("PRJ", db_manager=mock_db, user_role="supervisor")
        service = DashboardSyncService(manager)
        project = manager.get_project("PRJ")
        service.mirror_shots_to_excel(handler.read_shots(), project, handler, force=True)

        shot = next(s for s in handler.read_shots() if s.shot_name == "SH010")
        shot.sow = "remove wires"
        shot.dept("comp").artist = "priya"
        handler.write_shots([shot])
        service.mirror_shots_to_excel([shot], project, handler, force=True)

        rows = _rows_by_shot(_sheet(central / "Tracking" / "PRJ_tracking.xlsx"))

        assert len(rows) == 2
        assert rows["SH010"]["SOW"] == "remove wires"
        assert rows["SH010"]["COMP ARTIST"] == "priya"

    def test_a_shot_added_after_the_passbook_exists_still_gets_a_row(
            self, central, mock_db, tmp_path):
        manager = _ingested_project(mock_db, tmp_path)
        handler = SQLiteHandler("PRJ", db_manager=mock_db, user_role="supervisor")
        service = DashboardSyncService(manager)
        project = manager.get_project("PRJ")
        service.mirror_shots_to_excel(handler.read_shots(), project, handler, force=True)

        late = Shot(shot_name="SH030", reel_episode="ReelB")
        handler.write_shots([late])
        service.mirror_shots_to_excel([late], project, handler, force=True)

        rows = _rows_by_shot(_sheet(central / "Tracking" / "PRJ_tracking.xlsx"))

        assert "SH030" in rows

    def test_an_unreachable_central_folder_is_reported_not_swallowed(
            self, mock_db, tmp_path, monkeypatch):
        blocker = tmp_path / "not_a_folder"
        blocker.write_text("this file sits where the folder should be")
        monkeypatch.setattr(GlobalConfig, "server_root", classmethod(lambda cls: blocker))
        manager = _ingested_project(mock_db, tmp_path)
        handler = SQLiteHandler("PRJ", db_manager=mock_db, user_role="supervisor")
        service = DashboardSyncService(manager)

        ok, _ = service.mirror_shots_to_excel(
            handler.read_shots(), manager.get_project("PRJ"), handler, force=True)

        assert ok is False
        assert service.last_backup_error


class TestNothingReadsFromTheSheet:

    def test_the_service_has_no_import_at_all(self):
        assert not hasattr(DashboardSyncService, "sync_excel_to_database")
        assert not hasattr(DashboardSyncService, "is_excel_newer_than_db")
