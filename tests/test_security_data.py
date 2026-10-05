"""
2.2.0 hardening, data track: one check per rule.

The SQL console's read-only rule (SYS-002/016/041-044) is checked on both
databases in tests/test_write_results.py::test_execute_sql_is_read_only.
"""

import pytest


# ------------------------------------------------------------ SYS-020

def test_data_center_hides_secrets_and_keeps_users_read_only(pg_db, qtbot):
    from PySide6.QtCore import Qt
    from slate.core.infra.app_context import AppContext
    from slate.gui import database_explorer as dx
    from slate.core.infra.postgres_manager import PostgresManager
    PostgresManager._circuit_breaker.reset()

    ctx = AppContext(db_manager=pg_db)
    ctx.set_current_user({"username": "admin", "roles": ["Admin"]})
    view = dx.DatabaseExplorer(pg_db, app_context=ctx)
    qtbot.addWidget(view)
    data = {"cols_res": [{"column_name": c} for c in
                         ("id", "username", "password_hash", "api_token", "client_secret", "roles")],
            "rows": [{"id": 1, "username": "ana", "password_hash": "pbkdf2$x", "api_token": "t",
                      "client_secret": "s", "roles": '["Admin"]'}],
            "keys": ["id"]}
    view.show_table("ut_users", data)
    assert view.columns == ["id", "username", "roles"]
    assert view.key_columns == [] and not view.lbl_note.isHidden()
    roles = view.data_grid.item(0, view.columns.index("roles"))
    assert not (roles.flags() & Qt.ItemFlag.ItemIsEditable)
    view.apply_filter("pbkdf2")                      # search cannot find a hidden value
    assert view.data_grid.isRowHidden(0)

    view.txt_sql.setPlainText("SELECT 'h' AS password_hash, 1 AS one")
    view.run_custom_sql()
    qtbot.waitUntil(lambda: view.lbl_table_name.text().startswith("SQL Result"), timeout=5000)
    headers = [view.data_grid.horizontalHeaderItem(c).text() for c in range(view.data_grid.columnCount())]
    assert headers == ["one"]
    view.cancel_workers()


# ------------------------------------------------------------ DSH2-024 / SYS2-004

def test_neutralise_round_trips_exactly():
    from slate.core.domain.table_export import neutralise, restore
    for text in ("=HYPERLINK(\"x\")", "-5", "'=already", "''+x", "plain", "it's", ""):
        assert restore(neutralise(text)) == text
    assert neutralise("=1+1").startswith("'")


def test_excel_backup_is_neutralised_and_reads_back(tmp_path):
    from openpyxl import Workbook, load_workbook
    from slate.gui.tabs.vfx_dashboard_pro.core.excel_handler import ExcelHandler
    from slate.gui.tabs.vfx_dashboard_pro.core.project_manager import ProjectConfig
    from slate.gui.tabs.vfx_dashboard_pro.models.shot_model import FeedbackEntry, Shot

    wb = Workbook()
    ws = wb.active
    ws.title = "MASTER"
    ws.append(["Shot", "Reel", "Notes", "Client"])
    ws.append([""] * 4)
    path = tmp_path / "p.xlsx"
    wb.save(path)
    project = ProjectConfig(code="PRJ", name="P", project_number=1, excel_path=str(path),
                            sheet_name="MASTER", header_row=1, data_start_row=3,
                            column_mapping={"shot_name": "A", "reel": "B", "description": "C",
                                            "client_feedback": "D"})
    evil = '=HYPERLINK("http://x","y")'
    shot = Shot(shot_name="SH010", reel_episode="R01")
    shot.description = evil
    shot.feedback_client.append(FeedbackEntry(text="-fix " + evil, source="Client"))
    assert ExcelHandler(str(path), project).write_shots([shot])

    book = load_workbook(path)
    assert book["MASTER"]["C3"].value == "'" + evil
    assert book["MASTER"]["C3"].data_type == "s"
    assert all(not str(c.value or "").startswith(("=", "-")) for row in book["FEEDBACK_LOG"].iter_rows()
               for c in row)

    got = ExcelHandler(str(path), project).read_shots()[0]
    assert got.description == evil
    assert got.feedback_client[0].text == "-fix " + evil
    assert ExcelHandler(str(path), project).write_shots([got])      # a second save does not double up
    assert load_workbook(path)["MASTER"]["C3"].value == "'" + evil


def test_fleet_reports_are_neutralised(tmp_path):
    from openpyxl import load_workbook
    from slate.gui import admin_fleet_report_service as svc
    from slate.gui.admin_fleet_export import export_fleet_xlsx

    records = [{"pc_name": "=cmd|' /C calc'!A0", "status": "Online"}]
    summary = {"online": 1, "not_responding": 0, "offline": 0, "unknown": 0}
    svc.write_csv(tmp_path / "f.csv", records, summary, 0)
    assert "'=cmd" in (tmp_path / "f.csv").read_text(encoding="utf-8-sig")

    export_fleet_xlsx(str(tmp_path / "f.xlsx"), records, summary, 0,
                      columns=svc.columns_of(records), header_for=svc.header_for)
    values = [c.value for ws in load_workbook(tmp_path / "f.xlsx").worksheets
              for row in ws.iter_rows() for c in row if c.value]
    assert "'=cmd|' /C calc'!A0" in values and "=cmd|' /C calc'!A0" not in values


# ------------------------------------------------------------ SHL-120

def test_no_command_listener_is_started():
    import importlib.util
    import pathlib
    from slate.core.infra.network_manager import NetworkManager
    assert not hasattr(NetworkManager, "start") and not hasattr(NetworkManager, "_handle_tcp_client")
    assert importlib.util.find_spec("slate.gui.components.network_handler") is None
    import slate.gui.main_window as mw
    assert "NetworkManager" not in pathlib.Path(mw.__file__).read_text(encoding="utf-8")


# ------------------------------------------------------------ NEW-4

def test_a_saved_server_is_not_replaced_by_another_answer(monkeypatch):
    from slate.core.infra.global_config import GlobalConfig
    from slate.core.infra.postgres_manager import PostgresManager
    saved = {}
    monkeypatch.setattr(GlobalConfig, "set", classmethod(lambda cls, k, v: saved.__setitem__(k, v)))
    m = PostgresManager.__new__(PostgresManager)
    m.host_candidates, m.host, m.port, m.pooler_port = ["10.0.0.5"], "10.0.0.5", 5442, 6432

    m.saved_host = "10.0.0.5"
    assert m.adopt_announcement({"host": "10.0.0.66", "db_port": 5442, "pooler_port": 0}) is False
    assert m.host == "10.0.0.5" and saved == {} and "10.0.0.66" in m.discovery_warning
    assert m.adopt_announcement({"host": "10.0.0.5", "db_port": 5445, "pooler_port": 0}) is True

    m.saved_host = ""                                   # first setup: discovery still works
    m.host_candidates = ["127.0.0.1"]
    assert m.adopt_announcement({"host": "10.0.0.9", "db_port": 5442, "pooler_port": 0}) is True
    assert saved["db_host"] == "10.0.0.9"


# ------------------------------------------------------------ Paths

def test_built_paths_stay_inside_the_project(tmp_path):
    from slate.core.domain.naming import path_inside
    from slate.core.domain.shot_media import shot_folder

    assert path_inside(tmp_path, "05_Reels/R01", "SH010") == tmp_path / "05_Reels" / "R01" / "SH010"
    for bad in ("../x", "SH010/../../x", "C:/Windows" if tmp_path.drive != "C:" else "D:/x"):
        with pytest.raises(ValueError, match="outside the project folder"):
            path_inside(tmp_path, bad)

    class S:
        folder_paths = {"scan": "05_Reels/R01/../../../elsewhere", "comp": "05_Reels/R01/SH010/03_Comp"}
    assert shot_folder(S, tmp_path, "scan") is None
    assert shot_folder(S, tmp_path, "comp") == tmp_path / "05_Reels/R01/SH010/03_Comp"
